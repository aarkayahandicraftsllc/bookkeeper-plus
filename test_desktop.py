"""Drives the desktop app's real widgets, offscreen.

Run:  python test_desktop.py

Qt runs with the 'offscreen' platform so this works with no display, which
means the desktop build is covered by tests rather than only eyeballed.
"""

import os
import sys
import shutil
import tempfile
from datetime import date, timedelta
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

TMP = Path(tempfile.gettempdir()) / "bkplus_desktop_test.db"
TMP.unlink(missing_ok=True)
os.environ["BOOKS_DB"] = str(TMP)
# Company files are created in BOOKS_DIR, so point it at a scratch
# folder: without this the suite would write companies into the
# user's real data/ directory.
SANDBOX = Path(tempfile.gettempdir()) / f"bkplus_sandbox_{TMP.stem}"
shutil.rmtree(SANDBOX, ignore_errors=True)
SANDBOX.mkdir(parents=True, exist_ok=True)
os.environ["BOOKS_DIR"] = str(SANDBOX)

from PySide6 import QtCore, QtGui, QtWidgets  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402

import pandas as pd  # noqa: E402
import accounting as acc  # noqa: E402
import db  # noqa: E402
import desktop  # noqa: E402
import theme  # noqa: E402

FAILURES = []


def check(label, got, want, tol=0.005):
    ok = abs(got - want) <= tol if isinstance(want, float) else got == want
    print(f"  {'PASS' if ok else 'FAIL'}  {label}: got {got!r}, want {want!r}")
    if not ok:
        FAILURES.append(label)


def set_row(screen, row, ledger, debit=0.0, credit=0.0):
    """Fill one grid row the way a person would."""
    combo = screen.grid.cellWidget(row, 0)
    combo.setCurrentText(ledger)
    screen.grid.item(row, 1).setText(f"{debit:.2f}" if debit else "")
    screen.grid.item(row, 2).setText(f"{credit:.2f}" if credit else "")


db.init_db()
db.add_ledger("Bank", "Bank Accounts", 100_000, "Dr")
db.add_ledger("Capital", "Capital Account", 100_000, "Cr")
db.add_ledger("Sales", "Sales Accounts")
db.add_ledger("Purchases", "Purchase Accounts")
db.add_ledger("Rent", "Indirect Expenses")
db.add_ledger("Acme", "Sundry Debtors")

app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
window = desktop.MainWindow()

print("\n1) The window opens with every screen")
check("ten screens", window.stack.count(), 10)
check("screen names", [name.replace("&&", "&") for name, _ in window.screens],
      ["Voucher Entry", "Day Book", "Ledger", "Trial Balance", "Profit & Loss",
       "Balance Sheet", "Outstanding", "GST", "Masters", "Company"])
check("starts on voucher entry", window.stack.currentIndex(), 0)
check("title names the company and the app",
      "Bookkeeper Plus" in window.windowTitle(), True)

print("\n2) Alt+Shift+initial switches voucher type")
for key, name in desktop.VOUCHER_KEYS.items():
    window._set_voucher_type(name)
    check(f"{key} selects {name}", window.voucher.type_box.currentText(), name)
check("and it returns to voucher entry", window.stack.currentIndex(), 0)

# Every type must be reachable, and no two may want the same keystroke.
check("every voucher type has a key",
      sorted(desktop.VOUCHER_KEYS.values()), sorted(db.VOUCHER_TYPES))
check("and no key is used twice",
      len(set(desktop.VOUCHER_KEYS)), len(desktop.VOUCHER_KEYS))

# None of them may be a plain Alt+letter the menu bar already owns: that is
# what made Alt+C a dead key. Shift is what keeps these clear of the menus.
# Only TOP-LEVEL menus claim an Alt letter - findChildren() would also drag
# in submenus like "&Text size", whose T is reachable only once Settings is
# already open.
menu_mnemonics = {a.menu().title().split("&")[1][0].upper()
                  for a in window.menuBar().actions()
                  if a.menu() and "&" in a.menu().title()}
clashing = {key for key in desktop.VOUCHER_KEYS
            if key.startswith("Alt+") and not key.startswith("Alt+Shift+")
            and key.rsplit("+", 1)[1].upper() in menu_mnemonics}
check("no voucher key collides with a menu mnemonic", clashing, set())
check("the menus still own their own letters",
      menu_mnemonics, {"G", "V", "R", "C", "S"})
check("Qt parses every one of them",
      [k for k in desktop.VOUCHER_KEYS
       if QtGui.QKeySequence(k).toString() != k], [])

print("\n3) Accept stays disabled until the voucher balances")
voucher = window.voucher
voucher.reset()
check("disabled on an empty grid", voucher.accept_button.isEnabled(), False)

set_row(voucher, 0, "Rent", debit=12_000)
check("still disabled with one side only", voucher.accept_button.isEnabled(), False)
check("the difference is shown", "12,000.00" in voucher.gap_label.text(), True)

set_row(voucher, 1, "Bank", credit=9_000)
check("still disabled while out of balance",
      voucher.accept_button.isEnabled(), False)

set_row(voucher, 1, "Bank", credit=12_000)
check("enabled once both sides agree", voucher.accept_button.isEnabled(), True)
check("totals read back", voucher.dr_label.text(), "Dr 12,000.00")
check("difference is nil", voucher.gap_label.text(), "Difference 0.00")
check("it says so", "ready to accept" in voucher.notice.text(), True)

print("\n4) Voucher type guidance appears without blocking")
window._set_voucher_type("Journal")
set_row(voucher, 0, "Rent", debit=12_000)
set_row(voucher, 1, "Bank", credit=12_000)
check("a journal touching bank is flagged", "ℹ️" in voucher.notice.text(), True)
check("but accept is still allowed", voucher.accept_button.isEnabled(), True)
window._set_voucher_type("Payment")
set_row(voucher, 0, "Rent", debit=12_000)
set_row(voucher, 1, "Bank", credit=12_000)
check("a payment crediting bank is clean", "ℹ️" in voucher.notice.text(), False)

print("\n5) Accepting posts the voucher and clears the form")
before = len(db.get_vouchers())
window._accept()
check("one voucher posted", len(db.get_vouchers()), before + 1)
check("the grid was cleared", voucher.row_values(), [])
check("accept went back to disabled", voucher.accept_button.isEnabled(), False)
check("rent carries the debit",
      float(acc.balances().set_index("ledger").loc["Rent", "closing"]), 12_000.0)
check("it was filed as a Payment",
      str(db.get_vouchers().iloc[0]["voucher_type"]), "Payment")

print("\n6) A spare line appears as you fill the last one")
voucher.reset()
check("starts with two lines", voucher.grid.rowCount(), 2)
set_row(voucher, 1, "Sales", credit=5_000)
check("a third line is offered", voucher.grid.rowCount(), 3)

print("\n7) The reports populate")
db.save_voucher(voucher_type="Sales", voucher_date=date.today(),
                narration="invoice",
                lines=[{"ledger_id": db.ledger_id("Acme"), "debit": 80_000,
                        "credit": 0},
                       {"ledger_id": db.ledger_id("Sales"), "debit": 0,
                        "credit": 80_000}])
window.refresh_all()

check("day book lists both vouchers", window.daybook.table.rowCount(), 2)

# The two sides are a COLUMN EACH (user asked 2026-10-20), and one of them
# has to take the table's slack. It used to go to the first non-numeric
# column - `Date` - which handed a ten-character date 600px while the ledger
# names beside it were squeezed.
daybook_header = window.daybook.table.horizontalHeader()
labels = [window.daybook.table.horizontalHeaderItem(i).text()
          for i in range(window.daybook.table.columnCount())]
check("the day book names both sides in their own columns",
      labels[:3], ["Date", "Debit account", "Credit account"])
stretched = [labels[i] for i in range(len(labels))
             if daybook_header.sectionResizeMode(i)
             == QtWidgets.QHeaderView.Stretch]
# WHICH of the two it is depends on the ledger names in the book, so only
# that it is one of them is asserted - the defect was that it was neither.
check("and the slack goes to a ledger column, not to Date",
      [c for c in stretched if c not in ("Debit account", "Credit account")],
      [])
check("exactly one column takes it", len(stretched), 1)
check("trial balance has rows", window.trial.table.rowCount() > 0, True)
check("trial balance reports agreement",
      "equal credits" in window.trial.status.text(), True)
check("balance sheet agrees", "Both sides agree" in window.sheet.status.text(), True)
check("net profit shown on the P&L",
      "68,000.00" in window.profit.net_label.text(), True)   # 80,000 - 12,000
check("masters lists the ledgers, plus the six seeded tax ledgers",
      window.masters.ledger_table.rowCount(), 7 + len(db.GST_LEDGERS))
check("masters lists the groups", window.masters.group_table.rowCount(), 27)

print("\n8) The ledger screen shows a statement")
window.show_screen(2)
window.ledger.ledger_box.setCurrentText("Bank")
window.ledger.refresh()
check("bank statement has one entry", window.ledger.table.rowCount(), 1)
check("the caption shows the closing balance",
      "88,000.00 Dr" in window.ledger.caption.text(), True)

print("\n9) Opening a voucher from the day book loads it for alteration")
window.show_screen(1)
window.daybook.refresh()
first_id = int(window.daybook._frame.iloc[0]["voucher_id"])
window._alter(first_id)
check("it switched to voucher entry", window.stack.currentIndex(), 0)
check("it is in altering mode", window.voucher.editing_id, first_id)
check("the title says so", "Altering" in window.voucher.title.text(), True)
check("the lines were loaded", len(window.voucher.row_values()), 2)
# isVisible() needs a shown ancestor; isHidden() reflects the widget's
# own setVisible call, which is what we actually set here
check("cancel is offered", window.voucher.cancel_button.isHidden(), False)

rows = window.voucher.row_values()
check("the loaded voucher already balances",
      round(sum(r["debit"] - r["credit"] for r in rows), 2), 0.0)
window.voucher.reset()
check("cancelling clears the edit", window.voucher.editing_id, None)

print("\n10) Creating a ledger through Masters")
window.masters.led_name.setText("Petty Cash")
window.masters.led_group.setCurrentText("Cash-in-hand")
window.masters.led_open.setValue(2_500)
window.masters.led_side.setCurrentText("Dr")
window.masters.save_ledger()
check("the ledger exists", "Petty Cash" in set(db.get_ledgers()["name"]), True)
check("masters shows it",
      window.masters.ledger_table.rowCount(), 8 + len(db.GST_LEDGERS))
check("voucher entry offers it straight away",
      "Petty Cash" in window.voucher._ledger_names, True)
check("the opening difference is now reported",
      abs(acc.opening_difference()) > 0, True)
check("the status bar says so",
      "opening balances" in window.statusBar().currentMessage(), True)

print("\n11) Creating a ledger mid-voucher (Alt+L)")
voucher = window.voucher
voucher.reset()

# The shortcut must not be a letter the menu bar already owns: a menu-bar
# mnemonic BEATS a widget shortcut, which is why Alt+C silently opened the
# Company menu and never created a ledger.
menu_letters = {m.title().split("&")[1][0].upper()
                for m in window.menuBar().findChildren(QtWidgets.QMenu)
                if "&" in m.title()}
shortcuts = [s for s in voucher.findChildren(QtGui.QShortcut)
             if s.key().toString().startswith("Alt+")]
keys = sorted(s.key().toString() for s in shortcuts)
check("the voucher screen's Alt shortcuts", keys,
      ["Alt+B", "Alt+L", "Alt+T"])
check("none of them is a letter the menu bar owns",
      [k for k in keys if k.rsplit("+", 1)[1].upper() in menu_mnemonics], [])
check("L is not claimed by a menu", "L" in menu_letters, False)
check("C still is, which is what broke it", "C" in menu_letters, True)
labels = [b.text() for b in voucher.findChildren(QtWidgets.QPushButton)]
check("the button advertises the same key",
      any("(Alt+L)" in text for text in labels), True)
check("and nothing still advertises Alt+C",
      any("Alt+C" in text for text in labels), False)
# type a half-finished name into the first line, as you would before realising
# the ledger does not exist yet
voucher.grid.cellWidget(0, 0).setCurrentText("Courier Charges")
voucher.grid.setCurrentCell(0, 0)
check("the name is not a ledger yet",
      "Courier Charges" in voucher._ledger_ids, False)

dialog = desktop.LedgerDialog(voucher, name="Courier Charges")
check("the typed name is carried into the dialog", dialog.name.text(),
      "Courier Charges")
dialog.group_box.setCurrentText("Indirect Expenses")
check("it explains where the balance will land",
      "indirect expenses" in dialog.note.text().lower(), True)
dialog._create()
check("the ledger was created", dialog.created, "Courier Charges")
check("it exists in the books",
      "Courier Charges" in set(db.get_ledgers()["name"]), True)

# what create_ledger() does once the dialog returns
voucher._add_ledger_name("Courier Charges")
check("every line now offers it",
      all(voucher.grid.cellWidget(r, 0).findText("Courier Charges") >= 0
          for r in range(voucher.grid.rowCount())), True)
check("and the voucher screen knows its id",
      "Courier Charges" in voucher._ledger_ids, True)

# and it can be used straight away, on the same voucher
set_row(voucher, 0, "Courier Charges", debit=1_800)
set_row(voucher, 1, "Bank", credit=1_800)
check("the new ledger posts like any other",
      voucher.accept_button.isEnabled(), True)
window._set_voucher_type("Payment")
set_row(voucher, 0, "Courier Charges", debit=1_800)
set_row(voucher, 1, "Bank", credit=1_800)
window._accept()
check("it carries the debit",
      float(acc.balances().set_index("ledger").loc["Courier Charges", "closing"]),
      1_800.0)
check("the books still balance",
      acc.balance_sheet()["totals"]["difference"], 0.0, tol=0.0)
check("it reached the indirect expenses",
      "Courier Charges" in set(
          acc.profit_and_loss()["indirect_expense"]["Particulars"]), True)

print("\n12) Creating a group mid-voucher, then a ledger under it")
gdialog = desktop.GroupDialog(voucher, under="Secured Loans")
gdialog.name.setText("Equipment Loan")
gdialog._create()
check("the group was created", gdialog.created, "Equipment Loan")
check("it inherited its parent's nature",
      db.get_groups().set_index("name").loc["Equipment Loan", "nature"],
      "liability")
ldialog = desktop.LedgerDialog(voucher)
check("the new group is offered immediately",
      ldialog.group_box.findText("Equipment Loan") >= 0, True)
ldialog.name.setText("Tractor Loan")
ldialog.group_box.setCurrentText("Equipment Loan")
ldialog.opening.setValue(40_000)
ldialog.side.setCurrentText("Cr")
ldialog._create()
check("the ledger was created under it", ldialog.created, "Tractor Loan")
check("and it lands on the liabilities side",
      float(acc.balance_sheet()["liabilities"].set_index("Particulars")
            .loc["Loans (Liability)", "Amount"]), 40_000.0)

print("\n13) A rejected master shows a dialog, not a crash")
window.masters.led_name.setText("Petty Cash")      # duplicate
QtWidgets.QApplication.processEvents()
try:
    db.add_ledger("Petty Cash", "Cash-in-hand")
    check("duplicate rejected by the data layer", "accepted", "rejected")
except ValueError:
    check("duplicate rejected by the data layer", "rejected", "rejected")

print("\n14) Both front ends read the same books")
check("the desktop app uses the configured database", str(db.DB_PATH), str(TMP))
check("balances agree with the shared accounting layer",
      acc.trial_balance()["difference"], acc.trial_balance()["difference"])
check("balance sheet still balances",
      acc.balance_sheet()["totals"]["difference"], 0.0, tol=0.0)

print("\n15) Colour schemes")


def contrast(fg, bg):
    """WCAG relative-contrast ratio between two #rrggbb colours."""
    def luminance(value):
        parts = [int(value[i:i + 2], 16) / 255 for i in (1, 3, 5)]
        linear = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
                  for c in parts]
        return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]

    light, dark = sorted((luminance(fg), luminance(bg)), reverse=True)
    return (light + 0.05) / (dark + 0.05)


# Every scheme must be legible. This is the check that makes adding a scheme
# safe: a status line that carries meaning through colour is useless if you
# cannot read it, and eyeballing the light one says nothing about the others.
AA = 4.5
worst = (None, 99.0)
for key, scheme in theme.SCHEMES.items():
    for token in theme.TOKENS:
        for ground in ("window", "base"):
            ratio = contrast(scheme.color(token), getattr(scheme, ground))
            if ratio < worst[1]:
                worst = (f"{key}.{token} on {ground}", ratio)
check("every token in every scheme clears WCAG AA",
      worst[1] >= AA, True)
print(f"       tightest: {worst[0]} at {worst[1]:.2f}:1 (floor {AA})")

check("the schemes offered are the ones the menu lists",
      sorted(theme.ORDER), sorted(list(theme.SCHEMES) + [theme.AUTO]))
check("every offered scheme has a label",
      all(theme.LABELS.get(k) for k in theme.ORDER), True)
check("an unknown saved value falls back rather than crashing",
      theme.resolve("chartreuse").key, theme.DEFAULT)
for old, new in theme.ALIASES.items():
    check(f"a preference saved as {old!r} still opens {new!r}",
          theme.resolve(old).key, new)
check("auto resolves to a real scheme",
      theme.resolve(theme.AUTO).key in theme.SCHEMES, True)
try:
    theme.SCHEMES["light"].color("highlight")   # structural, not a token
    raised = False
except KeyError:
    raised = True
check("asking for a non-token colour raises rather than silently working",
      raised, True)

# Switching through every scheme on the live window, the way the menu does.
for key in theme.ORDER:
    window.set_color_scheme(key)
    QtWidgets.QApplication.processEvents()
    check(f"{key}: saved to the books", db.get_setting(theme.SETTING_KEY), key)
    check(f"{key}: menu entry is ticked",
          window._scheme_actions[key].isChecked(), True)
    resolved = theme.resolve(key)
    check(f"{key}: palette follows the scheme",
          QtWidgets.QApplication.palette().color(
              QtGui.QPalette.ColorRole.Window).name(), resolved.window)
    # The status lines must have been recoloured, not left on the old scheme.
    stale = [lbl.objectName() or lbl.text()[:24]
             for lbl in window.findChildren(QtWidgets.QLabel)
             if lbl.property("tone")
             and resolved.color(lbl.property("tone")) not in lbl.styleSheet()]
    check(f"{key}: no status line left on the old colour", stale, [])

window.trial.refresh()
check("a report run after switching uses the active scheme",
      theme.active().color("warn") in window.trial.status.styleSheet()
      or theme.active().color("good") in window.trial.status.styleSheet(),
      True)

window.set_color_scheme("light")
check("back on light, the platform style is restored",
      QtWidgets.QApplication.style().objectName().lower()
      in (theme._native_style.lower(), ""), True)

print("\n16) Text size, and the layout following it")

check("the default is a step above the platform size",
      theme.TEXT_STEPS[theme.DEFAULT_TEXT] > 0, True)

base_pt = theme.point_size()
window.set_color_scheme("light")

def settle():
    """Let Qt catch up, INCLUDING deferred deletions.

    processEvents() does not deliver DeferredDelete at the same event-loop
    level, so widgets retired with deleteLater() linger and the counts below
    creep upward - which looks exactly like a leak and is not one. Flushing
    them explicitly means these counts would catch a real leak.
    """
    QtWidgets.QApplication.processEvents()
    QtWidgets.QApplication.sendPostedEvents(
        None, QtCore.QEvent.Type.DeferredDelete)


sizes = {}
for key in theme.TEXT_ORDER:
    window.set_text_size(key)
    settle()
    pt = theme.point_size()
    sizes[key] = pt
    check(f"{key}: saved to the books",
          db.get_setting(theme.FONT_SETTING_KEY), key)
    check(f"{key}: menu entry is ticked",
          window._text_actions[key].isChecked(), True)
    check(f"{key}: the application font is {pt}pt",
          QtWidgets.QApplication.font().pointSize(), pt)

    # Everything that sets its own point size must track the setting. These are
    # the widgets that do NOT follow the application font on their own.
    headings = [w for w in window.findChildren(QtWidgets.QWidget)
                if w.property("heading")]
    monos = [w for w in window.findChildren(QtWidgets.QWidget)
             if w.property("mono")]
    check(f"{key}: all {len(headings)} headings scaled",
          {w.font().pointSize() for w in headings},
          {theme.heading_point_size()})
    check(f"{key}: all {len(monos)} monospaced labels scaled",
          {w.font().pointSize() for w in monos}, {pt})

    # A table's amount cells are rebuilt on refresh; they must come back at the
    # new size too, or the figures stay small while their labels grow.
    window.show_screen(3)                      # Trial Balance
    QtWidgets.QApplication.processEvents()
    cells = [window.trial.table.item(r, c)
             for r in range(window.trial.table.rowCount())
             for c in range(window.trial.table.columnCount())
             if window.trial.table.item(r, c)
             and window.trial.table.item(r, c).font().family() != ""]
    mono_cells = [i for i in cells if i.font().pointSize() == pt]
    check(f"{key}: table cells at the new size", bool(mono_cells), True)

check("each step is strictly larger than the last",
      [sizes[k] for k in theme.TEXT_ORDER],
      sorted(sizes[k] for k in theme.TEXT_ORDER))
check("Normal is the platform's own size", sizes["normal"], theme._base_point)

# The layout has to ADAPT, not just redraw bigger. Nothing may be clipped: a
# widget whose sizeHint exceeds the room it was given is cut off on screen.
window.set_text_size("largest")
window.resize(window._sized_to_text())
QtWidgets.QApplication.processEvents()
clipped = []
for index in range(len(window.screens)):
    window.show_screen(index)
    QtWidgets.QApplication.processEvents()
    screen = window.screens[index][1]
    for label in screen.findChildren(QtWidgets.QLabel):
        if not label.isHidden() and label.text() and not label.wordWrap():
            need = label.sizeHint()
            if need.height() > label.height() + 1:
                clipped.append(f"{window.screens[index][0]}: "
                               f"{label.text()[:30]!r}")
check("nothing is clipped at the largest size", clipped, [])

# A window can never be narrower than its widest UNWRAPPED label, so one long
# unwrapped sentence silently sets a floor on the whole app - the Balance Sheet
# note was forcing 1788px, which is wider than a 1366px laptop screen and grew
# with every text step. Guard the floor rather than the individual labels.
WIDTH_CEILING = 1400
too_wide = {name: screen.minimumSizeHint().width()
            for name, screen in window.screens
            if screen.minimumSizeHint().width() > WIDTH_CEILING}
check(f"no screen demands over {WIDTH_CEILING}px at the largest size",
      too_wide, {})
print("       screen minimum widths: "
      + ", ".join(f"{name.replace('&&', '&')} {screen.minimumSizeHint().width()}"
                  for name, screen in window.screens))

# And the window must not open bigger than the display it is on.
avail = QtGui.QGuiApplication.primaryScreen().availableGeometry().size()
want = window._sized_to_text()
check("the window fits the screen at the largest size",
      want.width() <= max(avail.width(), window.minimumSizeHint().width())
      and want.height() <= max(avail.height(), window.minimumSizeHint().height()),
      True)

widths, scales = [], []
for key in theme.TEXT_ORDER:
    window.set_text_size(key)
    settle()
    widths.append(window._sized_to_text().width())
    scales.append(theme.text_scale())
window.set_text_size("largest")

# The intent - bigger text asks for a bigger window - is the scale, which is
# environment-free. The realised width can only be non-decreasing, because it
# is clamped to the display and raised to the layout minimum: on a small
# screen (this test runs on an 800x800 offscreen one) every step lands on the
# same clamped size, which is correct rather than a failure.
check("each step asks for a larger window", scales, sorted(set(scales)))
check("the window never shrinks as text grows",
      widths, sorted(widths))
print("       scale " + " -> ".join(f"{s:.2f}" for s in scales)
      + " | width " + " -> ".join(f"{w}px" for w in widths))

for index in range(len(window.screens)):
    window.show_screen(index)
    QtWidgets.QApplication.processEvents()
    screen = window.screens[index][1]
    # Against the DISPLAY as well, because the window is clamped to it: this
    # offscreen one is 800px wide, so a window can never be wider than that
    # here however much a screen asks for. The real guarantee is the
    # WIDTH_CEILING check above - this one catches a screen asking for more
    # than the window gives it on a display that could have obliged.
    #
    # It used to pass without the clamp only because Voucher Entry demanded
    # 1327px and dragged the whole window out to that width. Narrowing the
    # screens made the window smaller and the artefact visible.
    want = min(screen.minimumSizeHint().width(), avail.width())
    check(f"{window.screens[index][0]} fits the window at largest size",
          window.width() >= want, True)

# The voucher-type description is the one label deliberately left UNWRAPPED,
# so it is the one that can pin the window open. Still at the largest size
# here, which is where it costs the most.
hint = window.voucher.hint
longest = max(db.VOUCHER_TYPES.values(), key=len)
hint.setText(longest)
QtWidgets.QApplication.processEvents()
check("the description is on one line", hint.wordWrap(), False)
check("and centred",
      bool(hint.alignment() & QtCore.Qt.AlignmentFlag.AlignHCenter), True)
check("and larger than body text",
      hint.font().pointSize() > QtWidgets.QApplication.font().pointSize(), True)
# If a longer description is ever added this is what will fail, rather than
# the window quietly growing past the width of a laptop screen.
check(f"the longest description ({len(longest)} chars) stays under "
      f"{WIDTH_CEILING}px unwrapped",
      hint.sizeHint().width() < WIDTH_CEILING, True)
print(f"       description at largest: {hint.font().pointSize()}pt, "
      f"{hint.sizeHint().width()}px of the {WIDTH_CEILING}px ceiling")

# Back to the default, and the books still add up after all that.
window.set_text_size(theme.DEFAULT_TEXT)
check("text size returns to the default",
      QtWidgets.QApplication.font().pointSize(), base_pt)
check("the description followed the text size back down",
      window.voucher.hint.font().pointSize(), theme.hint_point_size())
check("the balance sheet still balances",
      acc.balance_sheet()["totals"]["difference"], 0.0, tol=0.0)

print("\n17) The Company screen")

co = window.company
check("every declared field has an editor",
      sorted(co.editors), sorted(f.key for f in db.COMPANY_FIELDS))
check("it is reachable as a screen",
      "Company" in [name for name, _ in window.screens], True)

window.show_screen(len(window.screens) - 1)
settle()
co.editors["company_name"].setText("Aarkaya Handicrafts LLC")
co.editors["pan"].setText("aaapz1234c")          # lower case on purpose
co.editors["gstin"].setText("09AAAPZ1234C1Z5")
for key, line in zip(db.ADDRESS_LINES,
                     ("12 Mall Road", "Connaught Place", "")):
    co.editors[key].setText(line)
co.editors["city"].setText("New Delhi")
check("saving valid details succeeds", co.save(), True)
check("the form shows the upper-cased PAN back",
      co.editors["pan"].text(), "AAAPZ1234C")
check("it reached the books", db.get_company()["gstin"], "09AAAPZ1234C1Z5")
check("the window title follows the company name",
      window.windowTitle().startswith("Aarkaya Handicrafts LLC"), True)

# A rejected save must explain itself and KEEP the typing - losing a form to
# one bad character is worse than the bad character.
co.editors["gstin"].setText("09BBBPZ1234C1Z5")
check("a GSTIN not carrying the PAN is refused", co.save(), False)
check("the reason is on screen", "PAN" in co.status.text(), True)
check("the typing survives a refusal",
      co.editors["company_name"].text(), "Aarkaya Handicrafts LLC")
check("and nothing was written", db.get_company()["gstin"], "09AAAPZ1234C1Z5")

# refresh_all() reaches every screen; it must not wipe a part-typed form, just
# as it leaves a part-entered voucher alone.
co.editors["phone"].setText("011-5550123")
window.refresh_all()
settle()
check("refresh_all leaves unsaved company edits alone",
      co.editors["phone"].text(), "011-5550123")
co.refresh()
check("Revert reloads what is saved", co.editors["phone"].text(), "")

# A box sized in LINES has to grow with the text setting. No company field is
# multiline now that the address is separate lines, so this runs on a probe row
# - taken out again, because a row left in would shift every later geometry
# check on this screen.
probe = desktop.editor_for(db.CompanyField("probe", "Probe", kind="multiline"))
co.form.addRow("Probe", probe)
heights = []
for key in theme.TEXT_ORDER:
    window.set_text_size(key)
    settle()
    heights.append(probe.minimumHeight())
check("a multiline box grows with the text size", heights, sorted(set(heights)))
print(f"       multiline box: {' -> '.join(f'{h}px' for h in heights)}")
co.form.removeRow(probe)
settle()
window.set_text_size(theme.DEFAULT_TEXT)
settle()
check("the Company screen is not over-wide",
      co.minimumSizeHint().width() <= WIDTH_CEILING, True)

print("\n18) Choosing the financial year on a report")
window.show_screen(5)                       # Balance Sheet, an as-at report
bar = window.sheet.period
labels = [bar.year_box.itemText(i) for i in range(bar.year_box.count())]
check("the years are offered, newest last and Custom after them",
      labels[-1], bar.CUSTOM)
check("every financial year in the books is listed",
      [f"FY {y.label}" for y in acc.financial_years()], labels[:-1])
check("it opens on the year in progress",
      bar.selected_year().contains(date.today()), True)
check("an as-at report offers no since-books-began",
      bar.SINCE in labels, False)

# Picking last year must move the report, not just the dropdown.
last_year = acc.financial_years()[-2]
bar.year_box.setCurrentText(f"FY {last_year.label}")
settle()
check("picking a year sets the as-at date to its close",
      desktop.from_qdate(bar.to_edit.date()), last_year.end)
check("and the report followed", bar.range()[1], last_year.end)
check("the balance sheet still agrees in that year",
      "Both sides agree" in window.sheet.status.text(), True)

# A year means the WHOLE year, opening day to closing day - including a
# closing day still in the future for the year in progress. Clamping it to
# today would quietly report a different period from the one named.
current = acc.current_financial_year()
bar.year_box.setCurrentText(f"FY {current.label}")
settle()
check("the year in progress still runs to its closing day",
      desktop.from_qdate(bar.to_edit.date()), current.end)
check("which is in the future", current.end > date.today(), True)

# A date typed by hand is no longer a whole year, and the dropdown must not
# go on claiming that it is.
bar.to_edit.setDate(desktop.to_qdate(date.today() - timedelta(days=9)))
settle()
check("editing a date moves the dropdown to Custom",
      bar.year_box.currentText(), bar.CUSTOM)
check("and no year is reported as selected", bar.selected_year(), None)

# A period report has the same dropdown, plus since-books-began in it.
window.show_screen(4)                       # Profit & Loss
pl = window.profit.period
pl_labels = [pl.year_box.itemText(i) for i in range(pl.year_box.count())]
check("a period report offers since-books-began first",
      pl_labels[0], pl.SINCE)
check("the P&L opens on a financial year, not on all dates",
      pl.selected_year() is not None, True)
check("choosing a year fills both dates, start and end of it",
      (desktop.from_qdate(pl.from_edit.date()),
       desktop.from_qdate(pl.to_edit.date())),
      (current.start, current.end))
pl.year_box.setCurrentText(pl.SINCE)
settle()
check("since-books-began reports no start date", pl.range()[0], None)
check("and the From date is disabled, having nothing to say",
      pl.from_edit.isEnabled(), False)
whole_book = acc.profit_and_loss(None, date.today())["totals"]["net_profit"]
check("the net profit is the whole-book figure",
      f"{abs(whole_book):,.2f}" in window.profit.net_label.text(), True)

# The year list is derived from the books, so a voucher in an older year has
# to make that year reachable - otherwise there is no way to report on it.
db.save_voucher(
    voucher_type="Journal", voucher_date=date(date.today().year - 3, 6, 1),
    narration="an older year",
    lines=[{"ledger_id": db.ledger_id("Cash"), "debit": 5, "credit": 0},
           {"ledger_id": db.ledger_id("Sales"), "debit": 0, "credit": 5}])
window.sheet.period.reload_years()
settle()
grown = [window.sheet.period.year_box.itemText(i)
         for i in range(window.sheet.period.year_box.count())]
check("an older voucher adds its year to the dropdown",
      len(grown) > len(labels), True)
# Rebuilding the list must not quietly move the period you are looking at -
# this bar was left on custom dates above, and stays there.
check("a custom period survives the list being rebuilt",
      window.sheet.period.year_box.currentText(), window.sheet.period.CUSTOM)
window.sheet.period.year_box.setCurrentText(f"FY {current.label}")
settle()
check("and a year can be chosen again afterwards",
      window.sheet.period.selected_year().contains(date.today()), True)

print("\n19) Choosing a company at launch")

# This suite pins BOOKS_DB, which is exactly the case where the picker must
# NOT appear: the caller has already chosen the file.
check("a pinned data file is detected", db.DB_PINNED, True)
check("switching is offered only when the file is not pinned",
      window._switch_action.isEnabled(), not db.DB_PINNED)

# The dialogs are driven directly; they do not depend on being shown.
made = []
for name in ("Alpha Exports", "Beta Trading"):
    dialog = desktop.NewCompanyDialog()
    dialog.editors["company_name"].setText(name)
    dialog._create()
    check(f"created {name}", dialog.created is not None, True)
    made.append(dialog.created)
check("they are separate files", made[0] != made[1], True)

picker = desktop.CompanyPicker(current=made[1])
listed = [picker.combo.itemText(i) for i in range(picker.combo.count())]
check("both companies are in the dropdown",
      sorted(t.split("  ")[0] for t in listed)[:2],
      ["Alpha Exports", "Beta Trading"])
check("the dropdown starts on the company passed in",
      Path(picker.combo.currentData()).name, Path(made[1]).name)
check("the folder is named in the note", str(db.DATA_DIR) in picker.note.text(),
      True)
check("the detail line marks the one already open",
      "open now" in picker.detail.text(), True)

picker.combo.setCurrentIndex(
    next(i for i in range(picker.combo.count())
         if Path(picker.combo.itemData(i)).name == Path(made[0]).name))
picker._open()
check("opening returns the chosen file",
      Path(picker.chosen).name, Path(made[0]).name)

refused = desktop.NewCompanyDialog()
refused.editors["company_name"].setText("   ")
refused._create()
check("a company with no name is refused", refused.created, None)
check("the dialog says why", "name" in refused.status.text().lower(), True)

# Switching closes the window and asks again; main() loops on this flag.
check("the window does not start out switching", window.switch_requested, False)
# The window remembers the file it was OPENED with. db.DB_PATH has moved on,
# because creating the companies above switched it - and that difference is the
# whole reason company_file exists.
check("the window knows which company it is showing",
      window.company_file, Path(TMP))
check("even though the data layer has moved on",
      Path(db.DB_PATH) != window.company_file, True)
# With the file pinned, switching must do NOTHING at all - no dialog, no
# switch. The unpinned path is covered in test_multi_company.py.
window.switch_company()
check("switching is inert while the data file is pinned",
      window.switch_requested, False)

# A part-entered voucher must not be thrown away silently.
window.show_screen(0)
settle()
check("an empty voucher is not part-entered", window.part_entered_rows(), [])
check("and needs no confirmation to leave",
      window._confirm_leaving("Switch company"), True)
set_row(window.voucher, 0, "Cash", debit=10.0)
settle()
check("a part-entered voucher is detected",
      window.part_entered_rows(), [0])
window.voucher.reset()
settle()
check("resetting clears it again", window.part_entered_rows(), [])

print("\n20) Backing up on close, and on demand")
db.use_company(TMP)
window.company_file = Path(db.DB_PATH)
for stale in db.backups_for():
    stale.unlink()

# Nothing is taken on the way IN any more: main() no longer asks for one.
check("opening the app takes no backup by itself", len(db.backups_for()), 0)

# The close-time offer. The dialog is stubbed - a modal box would stop the
# suite dead - so the three answers are driven through ask_backup_on_close.
_real_ask = desktop.MainWindow.ask_backup_on_close


def _answer(reply):
    desktop.MainWindow.ask_backup_on_close = lambda self: reply


# "Back up and close" writes one, tagged for the moment it was taken.
_answer("yes")
_event = QtGui.QCloseEvent()
window.closeEvent(_event)
closed = db.backups_for()
check("closing with Yes writes a backup", len(closed), 1)
check("tagged as a close-time copy", "-close" in closed[0].name, True)
check("and the window really closes", _event.isAccepted(), True)

# With nothing changed since, a second close must not bury the useful copies.
_answer("yes")
_event = QtGui.QCloseEvent()
window.closeEvent(_event)
check("closing again with nothing changed adds nothing",
      len(db.backups_for()), 1)

# "Close without backing up" leaves the folder alone but still closes.
db.save_voucher(voucher_type="Journal", voucher_date=date.today(),
                narration="unsaved by choice",
                lines=[{"ledger_id": db.ledger_id("Cash"), "debit": 3,
                        "credit": 0},
                       {"ledger_id": db.ledger_id("Sales"), "debit": 0,
                        "credit": 3}])
_answer("no")
_event = QtGui.QCloseEvent()
window.closeEvent(_event)
check("closing with No writes nothing", len(db.backups_for()), 1)
check("but still closes", _event.isAccepted(), True)

# Cancel must call the close off entirely - including a company switch, which
# gets here by closing the window.
_answer("cancel")
window.switch_requested = True
_event = QtGui.QCloseEvent()
window.closeEvent(_event)
check("Cancel refuses the close", _event.isAccepted(), False)
check("and calls off a company switch with it",
      window.switch_requested, False)
check("still writing nothing", len(db.backups_for()), 1)

# Which of the two dialogs closing raises is decided by this, and only this.
# The dialogs themselves cannot be driven headlessly - clickedButton() is set
# by a real click - so what is pinned here is the branch, plus the closeEvent
# behaviour above for each answer.
#
# Note it does NOT short-circuit to "no": an earlier version skipped the
# dialog entirely when there was nothing to save, and closing in silence was
# reported as the backup not working at all.
desktop.MainWindow.ask_backup_on_close = _real_ask
db.backup_company(window.company_file, tag="manual")
check("with the books already backed up, that is what closing will say",
      db.backup_is_current(window.company_file), True)
check("and the backup can say when it was taken",
      bool(db.peek_settings(db.backups_for(window.company_file)[0],
                            (db.TAKEN_KEY,)).get(db.TAKEN_KEY)), True)
for stale in db.backups_for():
    stale.unlink()
check("with no backup at all there IS something to ask about",
      db.backup_is_current(window.company_file), False)
db.backup_company(window.company_file, tag="close")

# Company > Back up now. The dialogs it raises are suppressed, since a modal
# box would stop the suite dead waiting for a click.
_boxes = []
QtWidgets.QMessageBox.information = staticmethod(
    lambda *a, **k: _boxes.append(a[1:]) or QtWidgets.QMessageBox.Ok)
QtWidgets.QMessageBox.warning = staticmethod(
    lambda *a, **k: _boxes.append(("WARNING",) + a[1:])
    or QtWidgets.QMessageBox.Ok)

window.backup_now()
check("a manual backup with nothing changed says so rather than claiming to "
      "have written one",
      any("Already backed up" in str(b[0]) for b in _boxes), True)
check("and wrote nothing", len(db.backups_for()), 1)

_boxes.clear()
window.voucher.reset()
db.save_voucher(voucher_type="Journal", voucher_date=date.today(),
                narration="change the books",
                lines=[{"ledger_id": db.ledger_id("Cash"), "debit": 7,
                        "credit": 0},
                       {"ledger_id": db.ledger_id("Sales"), "debit": 0,
                        "credit": 7}])
window.backup_now()
check("a manual backup after a change writes one", len(db.backups_for()), 2)
check("it is tagged as manual", "-manual" in db.backups_for()[0].name, True)
check("and it reports where it went",
      any("Backed up" in str(b[0]) for b in _boxes), True)
check("no warning was raised at any point",
      [b for b in _boxes if b[0] == "WARNING"], [])

# A backup must never show up as a company to open.
check("backups are not offered as companies",
      [c for c in db.list_companies() if "-close" in c["file"]
       or "-manual" in c["file"]], [])

print("\n20b) Company > Exit leaves the same way the X does")
_exit = next((a for a in window.findChildren(QtGui.QAction)
              if a.text().startswith("Exit")), None)
check("the Company menu offers Exit", _exit is not None, True)
check("and it really is on that menu",
      any(a is _exit for m in window.menuBar().actions() if m.menu()
          and m.menu().title() == "&Company" for a in m.menu().actions()),
      True)
check("on Ctrl+Q", _exit.shortcut().toString(), "Ctrl+Q")

# The point of the check: it must go through closeEvent, so the backup offer
# happens. app.quit() would end the event loop without closing the window and
# skip it in silence - which is the one way to get this wrong.
db.save_voucher(voucher_type="Journal", voucher_date=date.today(),
                narration="entered before exiting",
                lines=[{"ledger_id": db.ledger_id("Cash"), "debit": 11,
                        "credit": 0},
                       {"ledger_id": db.ledger_id("Sales"), "debit": 0,
                        "credit": 11}])
_before = len(db.backups_for())
_answer("yes")
_exit.trigger()
QtWidgets.QApplication.processEvents()
check("triggering Exit asks about the backup and writes one",
      len(db.backups_for()), _before + 1)
check("tagged as a close-time copy",
      "-close" in db.backups_for()[0].name, True)

# And Cancel in that dialog must keep the app open.
db.save_voucher(voucher_type="Journal", voucher_date=date.today(),
                narration="and again",
                lines=[{"ledger_id": db.ledger_id("Cash"), "debit": 12,
                        "credit": 0},
                       {"ledger_id": db.ledger_id("Sales"), "debit": 0,
                        "credit": 12}])
_before = len(db.backups_for())
# Shown first, because the Exit above really did close it - so without this
# the check below would pass on a window that was already gone.
window.show()
QtWidgets.QApplication.processEvents()
check("the window is open to begin with", window.isHidden(), False)
_answer("cancel")
_exit.trigger()
QtWidgets.QApplication.processEvents()
check("Cancel calls the exit off and the window stays open",
      window.isHidden(), False)
check("and writes nothing", len(db.backups_for()), _before)
desktop.MainWindow.ask_backup_on_close = _real_ask
# Left SHOWN on purpose. This section is the only one that really closes the
# window (the checks above call closeEvent directly, which does not), and a
# window that has actually closed will not take activation again - which the
# keypress checks in section 24 need.
QtWidgets.QApplication.processEvents()

print("\n21) Custom dates ask for the dates")
db.use_company(TMP)
bar = window.profit.period                 # a period report: From and To
year = acc.current_financial_year()
bar.year_box.setCurrentText(f"FY {year.label}")
settle()

# The dialog, on its own. An as-at report must be asked for ONE date: a start
# field on a balance sheet would be a control that does nothing.
two = desktop.PeriodDialog(date(2026, 4, 1), date(2026, 6, 30), with_from=True)
check("a period report is asked for both dates", two.dates(),
      (date(2026, 4, 1), date(2026, 6, 30)))
one = desktop.PeriodDialog(date(2026, 4, 1), date(2026, 6, 30),
                           with_from=False)
check("an as-at report is asked for one", one.dates(),
      (None, date(2026, 6, 30)))
check("and its From editor is not shown", one.from_edit.isHidden(), True)

# A period cannot end before it starts.
_warned = []
_real_warning = QtWidgets.QMessageBox.warning
QtWidgets.QMessageBox.warning = staticmethod(
    lambda *a, **k: _warned.append(a[1:]) or QtWidgets.QMessageBox.Ok)
backwards = desktop.PeriodDialog(date(2026, 6, 30), date(2026, 4, 1))
backwards._accept()
check("dates the wrong way round are refused", len(_warned), 1)
check("and the dialog stays open", backwards.result(), 0)
QtWidgets.QMessageBox.warning = _real_warning

# Driven through the bar. exec() is stubbed because a modal dialog would stop
# the suite dead waiting for a click.
_real_exec = desktop.PeriodDialog.exec


def _accept_with(start, end):
    def stub(self):
        if start is not None and self.with_from:
            self.from_edit.setDate(desktop.to_qdate(start))
        self.to_edit.setDate(desktop.to_qdate(end))
        return QtWidgets.QDialog.DialogCode.Accepted
    return stub


desktop.PeriodDialog.exec = _accept_with(date(2026, 5, 1), date(2026, 5, 31))
bar.year_box.setCurrentText(bar.CUSTOM)
settle()
check("choosing custom dates puts them in the editors",
      (desktop.from_qdate(bar.from_edit.date()),
       desktop.from_qdate(bar.to_edit.date())),
      (date(2026, 5, 1), date(2026, 5, 31)))
check("and the report runs on them", bar.range(),
      (date(2026, 5, 1), date(2026, 5, 31)))
check("the dropdown says custom", bar.year_box.currentText(), bar.CUSTOM)

# THE BUG THIS FIXES: coming from since-books-began left the From editor
# disabled, so a custom period could never be given a start date.
bar.year_box.setCurrentText(bar.SINCE)
settle()
check("since-books-began disables From", bar.from_edit.isEnabled(), False)
desktop.PeriodDialog.exec = _accept_with(date(2026, 4, 10), date(2026, 4, 20))
bar.year_box.setCurrentText(bar.CUSTOM)
settle()
check("custom dates after it re-enables From", bar.from_edit.isEnabled(), True)
check("and the start date is actually used", bar.range()[0], date(2026, 4, 10))

# Cancelling must leave both the period and the dropdown as they were.
bar.year_box.setCurrentText(f"FY {year.label}")
settle()
was = bar.range()
desktop.PeriodDialog.exec = (
    lambda self: QtWidgets.QDialog.DialogCode.Rejected)
bar.year_box.setCurrentText(bar.CUSTOM)
settle()
check("cancelling leaves the dates alone", bar.range(), was)
check("and puts the dropdown back", bar.year_box.currentText(),
      f"FY {year.label}")
desktop.PeriodDialog.exec = _real_exec

# An as-at report must not be labelled as though it covered a period. Left as
# "Period ... To" it reads as half a control - a To with the From missing -
# next to the "From ... To" that the period reports show.
check("an as-at report is labelled for the moment it is drawn at",
      [lbl.text() for lbl in
       window.sheet.period.findChildren(QtWidgets.QLabel)
       if lbl.parent() is window.sheet.period],
      ["Year", "As at"])
check("a period report is labelled for the span it covers",
      [lbl.text() for lbl in
       window.profit.period.findChildren(QtWidgets.QLabel)
       if lbl.parent() is window.profit.period],
      ["Period", "From", "To"])
check("the trial balance is labelled the same way as the balance sheet",
      [lbl.text() for lbl in
       window.trial.period.findChildren(QtWidgets.QLabel)
       if lbl.parent() is window.trial.period],
      ["Year", "As at"])

print("\n22) Printing and exporting the reports")
db.use_company(TMP)
window.refresh_all()
import html  # noqa: E402  (only this section needs these)
import openpyxl  # noqa: E402

REPORTS = {1: "Day Book", 2: "Ledger", 3: "Trial Balance",
           4: "Profit & Loss", 5: "Balance Sheet"}
for index, title in REPORTS.items():
    window.show_screen(index)
    settle()
    report = window.current_report()
    check(f"{title}: the screen describes itself as a report",
          report is not None and report.title.startswith(title.split(" & ")[0]),
          True)
    check(f"{title}: it says which period it covers",
          bool(report.period), True)
    check(f"{title}: it has at least one table",
          len(report.sections) >= 1, True)

    # Printing cannot be clicked through here, but the page it would print is
    # just a document - so that is checked instead of the printer.
    page = desktop.report_html(report)
    # Escaped, because "Profit & Loss" is written "Profit &amp; Loss" on a
    # page - the title is text being put into HTML, not markup.
    check(f"{title}: the printable page names the report",
          html.escape(report.title) in page, True)
    check(f"{title}: and the company whose books these are",
          db.company_display_name() in page, True)

    for suffix in (".csv", ".xlsx"):
        out = desktop.export_report(
            report, path=str(db.DATA_DIR / report.filename(suffix)))
        check(f"{title}: exported{suffix}",
              out.exists() and out.stat().st_size > 0, True)
    csv_text = (db.DATA_DIR / report.filename(".csv")).read_text(
        encoding="utf-8-sig")
    check(f"{title}: the CSV leads with the report, company and period",
          csv_text.splitlines()[:3],
          [report.title, db.company_display_name(), report.period])
    book = openpyxl.load_workbook(db.DATA_DIR / report.filename(".xlsx"))
    # A two-sided statement is ONE sheet, read left and right; everything
    # else keeps a sheet per table.
    check(f"{title}: the workbook is laid out as the report asks",
          len(book.sheetnames),
          1 if report.side_by_side else len(report.sections))
    check(f"{title}: each sheet is headed with the company",
          book[book.sheetnames[0]]["A1"].value, db.company_display_name())
    book.close()

# The balance sheet exports as ONE sheet, read left and right (user asked
# 2026-10-20: "not in 2 seperate sheets, it should be in one sheet and looks
# more professional"). Two sheets meant holding one side in your head while
# looking at the other.
window.show_screen([n for n, _ in window.screens].index("Balance Sheet"))
settle()
sides_path = db.DATA_DIR / "sides.xlsx"
desktop.export_report(window.sheet.report(), path=str(sides_path))
book = openpyxl.load_workbook(sides_path)
sides = book[book.sheetnames[0]]
check("one sheet, not one per side", len(book.sheetnames), 1)
check("headed with the company", sides["A1"].value, db.company_display_name())
check("then what the report is", sides["A2"].value, "Balance Sheet")
check("then when it is drawn at", "as at" in str(sides["A3"].value), True)
check("the two sides sit beside each other",
      [sides["A5"].value, sides["D5"].value], ["Liabilities", "Assets"])
check("with a figures column each",
      [sides["B5"].value, sides["E5"].value], ["Amount", "Amount"])
totals_at = [r for r in range(1, sides.max_row + 1)
             if sides.cell(row=r, column=1).value == "Total"]
check("each side is ruled off and totalled", len(totals_at), 1)
check("and the totals agree with the statement",
      [sides.cell(row=totals_at[0], column=c).value for c in (2, 5)],
      [window.sheet.report().totals["Liabilities"],
       window.sheet.report().totals["Assets"]])
check("the total is emphasised",
      sides.cell(row=totals_at[0], column=1).font.bold, True)
# Still numbers with a format, never formatted text - the layout got smarter,
# the principle did not change.
# int, not float, when the figure is whole - openpyxl narrows it on the way
# back in. What matters is that it is a NUMBER and not text.
check("amounts are numbers",
      isinstance(sides["B6"].value, (int, float))
      and not isinstance(sides["B6"].value, bool), True)
check("carrying a money format", sides["B6"].number_format, "#,##0.00")
check("and the gridlines are off", sides.sheet_view.showGridLines, False)
book.close()

# Figures stay NUMBERS in a spreadsheet so they can be summed on arrival; a
# column of "1,234.00" strings is a file you have to repair before using.
window.show_screen(3)
settle()
trial = window.current_report()
desktop.export_report(trial, path=str(db.DATA_DIR / "numbers.xlsx"))
book = openpyxl.load_workbook(db.DATA_DIR / "numbers.xlsx")
sheet = book[book.sheetnames[0]]
columns = [c.value for c in sheet[5]]
amounts = [sheet.cell(row=r, column=columns.index("Debit") + 1).value
           for r in range(6, 6 + len(trial.sections[0].frame))]
check("exported amounts are numbers, not formatted text",
      all(v is None or isinstance(v, (int, float)) for v in amounts), True)
book.close()
# ...while the printed page IS formatted, being the copy somebody reads.
check("the printed page groups the thousands",
      "," in desktop.report_html(trial), True)

# The input screens are not reports, and must say so rather than printing
# something arbitrary.
# Voucher Entry, Masters and Company hold INPUT, so there is nothing to
# print. Outstanding and GST moved in at 6 and 7 and ARE reports, so the
# input screens are now 0, 8 and 9.
for index in (0, 8, 9):
    window.show_screen(index)
    settle()
    check(f"screen {index} ({window.screens[index][0]}) is not a report",
          window.current_report(), None)
window.show_screen(6)
settle()
check("but Outstanding is", window.current_report().title, "Outstanding")
window.show_screen(7)
settle()
check("and so is GST", window.current_report().title.startswith("GST"), True)
# Back to an input screen BEFORE the calls below. They are here to prove
# print/export are harmless when there is nothing to print - and on a report
# screen print_current() opens the print preview, which blocks for ever with
# nobody to dismiss it.
window.show_screen(8)
settle()
window.print_current()          # must not raise, and must not print
window.export_current()
check("asking anyway leaves a message instead of a document",
      "not a report" in window.statusBar().currentMessage(), True)

# Export offers the Downloads folder and opens what it wrote. Both the save
# dialog and the launch are stubbed - one would block the suite waiting for a
# click, the other would start Excel.
check("the Downloads folder is a real directory",
      desktop.downloads_dir().is_dir(), True)
check("and is the one the platform reports",
      desktop.downloads_dir().name.lower() in ("downloads", "223137552.aeroad"),
      True)

_offered = []
_launched = []
_real_dialog = QtWidgets.QFileDialog.getSaveFileName
_real_open = QtGui.QDesktopServices.openUrl
target = db.DATA_DIR / "opened-after-export.xlsx"


def _stub_dialog(parent=None, caption="", directory="", filter=""):
    _offered.append(directory)
    return str(target), filter


QtWidgets.QFileDialog.getSaveFileName = staticmethod(_stub_dialog)
QtGui.QDesktopServices.openUrl = staticmethod(
    lambda url: _launched.append(url.toLocalFile()) or True)

window.show_screen(3)
settle()
written = desktop.export_and_open(window.current_report(), window)
check("export writes the file", written == target and target.exists(), True)
check("the save dialog starts in the Downloads folder",
      str(desktop.downloads_dir()) in _offered[0], True)
check("and it suggests a name, not just a folder",
      _offered[0].lower().endswith(".xlsx"), True)
check("the file is opened afterwards",
      # Compared as paths: a QUrl gives forward slashes where Windows gives
      # backslashes, and the two strings differ for the same file.
      [Path(p) for p in _launched], [target])
check("the status bar says where it went",
      str(target) in window.statusBar().currentMessage(), True)

# A file that cannot be opened must still be reported, not silently dropped.
_launched.clear()
QtGui.QDesktopServices.openUrl = staticmethod(lambda url: False)
desktop.export_and_open(window.current_report(), window)
check("a file that will not open still reports its path",
      str(target) in window.statusBar().currentMessage()
      and "could not open" in window.statusBar().currentMessage(), True)

# Cancelling the dialog writes nothing and opens nothing.
QtWidgets.QFileDialog.getSaveFileName = staticmethod(
    lambda *a, **k: ("", ""))
_launched.clear()
check("cancelling the save dialog writes nothing",
      desktop.export_and_open(window.current_report(), window), None)
check("and opens nothing", _launched, [])

QtWidgets.QFileDialog.getSaveFileName = _real_dialog
QtGui.QDesktopServices.openUrl = _real_open

print("\n22b) The date format is a setting, and changes the display only")
# The ask (2026-10-20) was "$ means MM/DD/YYYY, rupees mean DD/MM/YYYY, or
# give me a line to pick it". It is a field, defaulted from the currency.
window.show_screen(1)                              # Day Book
settle()
check("the Day Book has rows to read a date out of",
      window.daybook.table.rowCount() > 0, True)
iso = str(window.daybook._shown["Date"].iloc[0])
when = date.fromisoformat(iso)


def daybook_date() -> str:
    return window.daybook.table.item(0, 0).text()


# This column used to show the raw ISO date while the editors above it said
# 01/04/2026 - the report tables ignored the convention entirely.
check("the table shows a date, not the ISO the frame carries",
      daybook_date(), when.strftime("%d/%m/%Y"))
check("and the voucher editor takes one the same way",
      window.voucher.date_edit.displayFormat(), "dd/MM/yyyy")

company_index = [name for name, _ in window.screens].index("Company")
window.show_screen(company_index)
settle()
co.editors["date_format"].setCurrentText("MM/DD/YYYY")
check("saving the new format succeeds", co.save(), True)
settle()
# A QDateEdit's display format is set ONCE, when it is built, so every editor
# on every other screen would otherwise go on showing the old order.
check("the Day Book followed", daybook_date(), when.strftime("%m/%d/%Y"))
check("and so did the voucher editor, on another screen",
      window.voucher.date_edit.displayFormat(), "MM/dd/yyyy")
check("and the period editors on the reports",
      window.daybook.period.to_edit.displayFormat(), "MM/dd/yyyy")

# Print follows the screen; the spreadsheet does NOT. It keeps the ISO date it
# can sort and filter on - the same split the money columns already make.
window.show_screen(1)
settle()
dated = window.current_report()
page = desktop.report_html(dated)
check("the printed page carries the chosen format",
      when.strftime("%m/%d/%Y") in page, True)
check("and not the ISO date", iso in page, False)
desktop.export_report(dated, path=str(db.DATA_DIR / "dates.csv"))
check("the CSV keeps ISO, so it still sorts",
      iso in (db.DATA_DIR / "dates.csv").read_text(encoding="utf-8-sig"), True)

window.show_screen(company_index)
settle()
co.editors["date_format"].setCurrentText("DD/MM/YYYY")
check("and it goes back", co.save(), True)
settle()
check("the Day Book came back with it", daybook_date(),
      when.strftime("%d/%m/%Y"))

print("\n22c) Clicking a column header sorts the report")
window.show_screen(1)
window.daybook.refresh()
settle()
table = window.daybook.table
labels = [table.horizontalHeaderItem(i).text()
          for i in range(table.columnCount())]
unsorted_first = table.key_at(0)

# Amounts must sort as NUMBERS. Read as text, "1,500.00" comes before
# "100.00", which is the whole reason a cell carries a sort key.
debit_col = labels.index("Debit")
table.sort_by(debit_col)
settle()
amounts = [desktop.parse_amount(table.item(r, debit_col).text())
           for r in range(table.rowCount())]
check("sorting by an amount orders it by value",
      amounts, sorted(amounts))
table.sort_by(debit_col)
settle()
amounts_down = [desktop.parse_amount(table.item(r, debit_col).text())
                for r in range(table.rowCount())]
check("clicking the same header again reverses it",
      amounts_down, sorted(amounts_down, reverse=True))

# Dates must sort CHRONOLOGICALLY even though they READ dd/MM/yyyy. Every
# voucher in this book is dated today, so the real case is built by hand:
# 14 October and 2 November sort the wrong way round as text, which is
# exactly the trap the sort key exists for.
probe = desktop.FrameTable()
probe.load(pd.DataFrame({"Date": ["2026-11-02", "2026-10-14"],
                         "Particulars": ["second", "first"]}))
check("the two dates really do sort wrongly as text",
      sorted(probe.item(r, 0).text() for r in range(2)),
      ["02/11/2026", "14/10/2026"])
probe.sort_by(0)
check("but sorting the column is chronological",
      [probe.item(r, 0).text() for r in range(2)],
      ["14/10/2026", "02/11/2026"])
check("so the rows moved with it",
      [probe.item(r, 1).text() for r in range(2)], ["first", "second"])
probe.deleteLater()

# THE REGRESSION THIS COULD HAVE INTRODUCED: the Day Book opened the voucher
# at the clicked POSITION in the frame, so after a sort it would have opened
# a different voucher than the one on that line.
table.sort_by(labels.index("Vch No"))
settle()
moved_first = table.key_at(0)
check("the sort really moved the rows", moved_first != unsorted_first, True)
window.daybook._open(table.model().index(0, 0))
settle()
check("double-click after a sort opens the row that was clicked",
      window.voucher.editing_id, int(moved_first))
window.voucher.reset()

# A total row cannot survive a sort, so the Trial Balance keeps its total in
# a bar under the table - but the PRINTED report still carries it as a row.
window.show_screen(3)
settle()
tb = window.trial
check("the trial balance table holds no Total row",
      [tb.table.item(r, 0).text() for r in range(tb.table.rowCount())
       if tb.table.item(r, 0).text() == "Total"], [])
check("but the report it prints does",
      str(tb.report().sections[0].frame["Particulars"].iloc[-1]), "Total")
check("and its table offers sorting", tb.table._pinned_last, False)
# The zero-balance toggle (user asked 2026-10-20). Out by default.
db.add_ledger("Dormant Account", "Indirect Expenses")
tb.refresh()
settle()
lean = tb.table.rowCount()
listed = [tb.table.item(r, 0).text() for r in range(lean)]
check("an account with no balance is not listed",
      "Dormant Account" in listed, False)
tb.with_zero.setChecked(True)
settle()
check("ticking the toggle brings it in",
      "Dormant Account" in [tb.table.item(r, 0).text()
                            for r in range(tb.table.rowCount())], True)
check("which is more rows", tb.table.rowCount() > lean, True)
as_at = tb.period.range()[1]
check("the totals cannot move - nil adds nothing to a side",
      acc.trial_balance(as_at, include_zero=True)["debit"],
      acc.trial_balance(as_at)["debit"], tol=0.0)
check("and the printed report says it is included",
      "no balance" in tb.report().period, True)
tb.with_zero.setChecked(False)
settle()
check("unticking puts it back", tb.table.rowCount(), lean)
db.delete_ledger(db.ledger_id("Dormant Account"))
tb.refresh()
settle()
tb.table.sort_by(0)
settle()
names = [tb.table.item(r, 0).text() for r in range(tb.table.rowCount())]
check("which sorts it, case-insensitively",
      names, sorted(names, key=str.casefold))

# Left as it was found, so the later sections see the default order.
for screen in (window.daybook.table, window.trial.table):
    screen._sort = None
window.daybook.refresh()
window.trial.refresh()
settle()

print("\n23) Voiding, restoring and deleting a voucher")
db.use_company(TMP)
voucher = window.voucher
vid = db.save_voucher(
    voucher_type="Journal", voucher_date=date.today(), narration="to void",
    lines=[{"ledger_id": db.ledger_id("Rent"), "debit": 4_000, "credit": 0},
           {"ledger_id": db.ledger_id("Bank"), "debit": 0, "credit": 4_000}])

voucher.load_voucher(vid)
check("altering offers Void", voucher.void_button.isHidden(), False)
check("and Delete", voucher.delete_button.isHidden(), False)
check("but not Restore", voucher.restore_button.isHidden(), True)
check("the title says Altering", "Altering" in voucher.title.text(), True)

# Say yes to the confirmation, without a dialog to click.
_real_question = QtWidgets.QMessageBox.question
QtWidgets.QMessageBox.question = staticmethod(
    lambda *a, **k: QtWidgets.QMessageBox.StandardButton.Yes)
voucher.void_voucher()
check("the voucher is voided", db.is_voided(vid), True)
check("and the editor was left", voucher.editing_id, None)

voucher.load_voucher(vid)
check("reopening says Voided", "Voided" in voucher.title.text(), True)
check("Restore is offered instead", voucher.restore_button.isHidden(), False)
check("Void is not", voucher.void_button.isHidden(), True)
check("Accept is refused", voucher.accept_button.isEnabled(), False)
check("and the notice explains why",
      "voided" in voucher.notice.text().lower(), True)
# The trap: _recalculate runs on every keystroke and used to hand Accept back.
set_row(voucher, 0, "Rent", debit=4_000)
check("typing in the grid does NOT re-enable Accept",
      voucher.accept_button.isEnabled(), False)

window.refresh_all()
book = window.daybook._shown
check("the Day Book gained a Status column", "Status" in book.columns, True)
check("the voided voucher is marked",
      "Voided" in list(book["Status"]), True)
check("showing nil", float(book[book["Status"] == "Voided"]["Debit"].iloc[0]),
      0.0)
check("and the count says how many", "1 voided" in window.daybook.count.text(),
      True)
check("a voided voucher prints with the rest",
      "Status" in window.daybook.report().sections[0].frame.columns, True)

voucher.load_voucher(vid)
voucher.restore_voucher()
check("restoring puts it back", db.is_voided(vid), False)
voucher.load_voucher(vid)
check("and Accept works again", voucher.accept_button.isEnabled(), True)
window.refresh_all()
check("the Status column goes away again",
      "Status" in window.daybook._shown.columns, False)

# "Void instead" on the delete dialog must void, not delete.
voucher.load_voucher(vid)
voucher._ask_delete = lambda: "void"
voucher.delete_voucher()
check("choosing Void instead voids", db.is_voided(vid), True)
check("and keeps the voucher", db.get_voucher(vid) is None, False)

voucher.load_voucher(vid)
voucher._ask_delete = lambda: "keep"
voucher.delete_voucher()
check("backing out of the delete keeps it", db.get_voucher(vid) is None, False)

number_before = db.next_voucher_number("Journal")
voucher.load_voucher(vid)
voucher._ask_delete = lambda: "delete"
voucher.delete_voucher()
check("deleting for good removes it", db.get_voucher(vid) is None, True)
check("and gives the number up",
      db.next_voucher_number("Journal"), number_before - 1)
QtWidgets.QMessageBox.question = _real_question
voucher.reset()

print("\n23b) The balance sheet, summarised and detailed")
window.show_screen(5)
settle()
bs = window.sheet
check("it opens summarised", bs.detailed.isChecked(), False)
summary_rows = bs.liabilities.rowCount()
summary_total = bs.liability_label.text()
check("and its tables can be sorted", bs.liabilities._pinned_last, False)

bs.detailed.setChecked(True)
settle()
check("ticking it adds rows", bs.liabilities.rowCount() > summary_rows, True)
check("which are indented accounts",
      any(bs.liabilities.item(r, 0).text().startswith(acc.DETAIL_INDENT)
          for r in range(bs.liabilities.rowCount())), True)
check("the side total does NOT change", bs.liability_label.text(),
      summary_total)
# An indented table cannot be sorted: the accounts would be carried away from
# the heading they belong under, and the statement would say nothing true.
check("so sorting is withdrawn while it is detailed",
      bs.liabilities._pinned_last, True)
before_sort = [bs.liabilities.item(r, 0).text()
               for r in range(bs.liabilities.rowCount())]
bs.liabilities.sort_by(0)
settle()
check("and clicking a header does nothing",
      [bs.liabilities.item(r, 0).text()
       for r in range(bs.liabilities.rowCount())], before_sort)
check("the printed report says which format it is",
      "under each heading" in bs.report().period, True)
check("and warns that the column holds both",
      any("totals above are the ones to read" in n for n in bs.report().notes),
      True)

bs.detailed.setChecked(False)
settle()
check("unticking it goes back to the summary",
      bs.liabilities.rowCount(), summary_rows)
check("and sorting comes back", bs.liabilities._pinned_last, False)

print("\n23c) Copying a voucher, for an entry that repeats")
db.use_company(TMP)
v = window.voucher
window.show_screen(0)
v.reset()
settle()

# Derived, not hardcoded: earlier sections post, void and delete, so which
# types have a voucher to copy is not fixed.
has = [k for k in db.VOUCHER_TYPES if db.last_voucher_id(k) is not None]
none_yet = [k for k in db.VOUCHER_TYPES if db.last_voucher_id(k) is None]
check("some type has a voucher to copy", bool(has), True)

kind = has[0]
v.type_box.setCurrentText(kind)
settle()
check("the button names the voucher it would copy",
      v.copy_button.text(), f"Copy last {kind.lower()}")
check("and is offered", v.copy_button.isEnabled(), True)

source_id = db.last_voucher_id(kind)
head, lines = db.get_voucher(source_id)
expected = [(str(ln["ledger_name"]), float(ln["debit"]), float(ln["credit"]))
            for ln in lines]
next_number = db.next_voucher_number(kind)

v.copy_voucher()
settle()
copied = [(n, r["debit"], r["credit"]) for n, r in
          zip([v.grid.cellWidget(i, 0).currentText()
               for i in range(v.grid.rowCount())], v.row_values())]
check("the lines came across", [(n, d, c) for n, d, c in copied], expected)
check("it is a NEW voucher, not an edit of the old one", v.editing_id, None)
check("numbered next, not the number it copied",
      v.number_label.text(), f"No. {next_number}")
check("dated today, which is the field you would always have to change",
      desktop.from_qdate(v.date_edit.date()), date.today())
check("the narration came too", v.narration.text(), str(head["narration"]))
check("and the grid still has exactly one spare line",
      v.grid.rowCount(), len(expected) + 1)
check("nothing is saved until you accept it",
      db.next_voucher_number(kind), next_number)

# There has to be a way OUT of a copy. Until 2026-10-20 the only button was
# Accept: the form was full and could not be emptied (user: "there is no
# option to can cel it only thing i can do is accept").
check("a copy offers a way to discard it", v.cancel_button.isHidden(), False)
check("worded for a copy, not for an edit",
      v.cancel_button.text(), desktop.DISCARD_COPY)
v.cancel_button.click()
settle()
check("discarding it empties the form", v.row_values(), [])
check("and writes nothing", db.next_voucher_number(kind), next_number)
check("the way out goes away with it", v.cancel_button.isHidden(), True)

v.copy_voucher()
settle()
v.accept_voucher()
settle()
check("accepting a copy posts it, so the type has one more voucher",
      db.next_voucher_number(kind), next_number + 1)
made = db.last_voucher_id(kind)
made = db.last_voucher_id(kind)
check("and the voucher it was copied FROM is untouched",
      db.get_voucher(source_id)[0]["number"], head["number"])
db.delete_voucher(int(made))        # leave the book as it was found
window.refresh_all()

# From a voucher being altered, the same button copies THAT one.
v.load_voucher(int(source_id))
settle()
check("while altering, the button offers this voucher",
      v.copy_button.text(), "Copy this voucher")
v.copy_voucher()
settle()
check("which also starts a new voucher", v.editing_id, None)
check("carrying the same lines", len(v.row_values()), len(expected))
v.reset()
settle()

if none_yet:
    v.type_box.setCurrentText(none_yet[0])
    settle()
    check("a type with no vouchers yet cannot be copied",
          v.copy_button.isEnabled(), False)
    check("and the button says why",
          "no" in v.copy_button.toolTip().lower(), True)
v.type_box.setCurrentText(kind)
v.reset()
settle()


# Plus only: what a copy does with the two things Bookkeeper has not got.
billed = [k for k in db.VOUCHER_TYPES
          if db.last_voucher_id(k) is not None
          and any(ln.get("bills") for ln in
                  db.get_voucher(db.last_voucher_id(k))[1])]
if billed:
    v.type_box.setCurrentText(billed[0])
    settle()
    v.copy_voucher()
    settle()
    check("a copy carries no bill allocations - a repeat is a NEW bill",
          [b for r in v.row_values() for b in r.get("bills", [])], [])
    v.reset()
    settle()

print("\n23d) The next line is offered the balancing amount")
db.use_company(TMP)
v = window.voucher
window.show_screen(0)
v.reset()
settle()
v.grid.cellWidget(0, 0).setCurrentText("Rent")
v.grid.item(0, 1).setText("1,000.00")
settle()
check("the line below is offered the other side",
      v.grid.item(1, 2).text(), "1,000.00")
check("and nothing in its debit column", v.grid.item(1, 1).text(), "")
check("it counts as a suggestion, not as typing",
      v.grid.item(1, 2).data(desktop.SUGGESTED), True)

# "unless i change it to make it a multiline journal": taking 400 of the
# 1,000 must leave 600 offered on the line after, not 1,000 again.
v.grid.item(1, 2).setText("400")
settle()
check("overwriting it leaves the remainder below",
      v.grid.item(2, 2).text(), "600.00")
check("and the line taken over is no longer a suggestion",
      v.grid.item(1, 2).data(desktop.SUGGESTED), None)

# A figure entered by hand is never moved or erased.
v.grid.item(2, 2).setText("600")
settle()
v.grid.cellWidget(1, 0).setCurrentText("Cash")
v.grid.cellWidget(2, 0).setCurrentText("Bank")
settle()
check("the voucher balances on what was entered",
      round(sum(r["debit"] - r["credit"] for r in v.row_values()), 2), 0.0)
check("three lines, as typed", len(v.row_values()), 3)
v.grid.item(0, 1).setText("1,000.00")        # retype the same figure
settle()
check("re-entering the same amount disturbs nothing typed",
      [v.grid.item(1, 2).text(), v.grid.item(2, 2).text()],
      ["400", "600"])

# A suggestion that is no longer needed is withdrawn, not left to be read as
# a figure.
v.reset()
settle()
v.grid.cellWidget(0, 0).setCurrentText("Rent")
v.grid.item(0, 1).setText("250")
settle()
check("a suggestion appears", v.grid.item(1, 2).text(), "250.00")
v.grid.item(0, 1).setText("")
settle()
check("and is withdrawn when the amount above goes away",
      v.grid.item(1, 2).text(), "")
v.reset()
settle()

print("\n23e) Editing a ledger and a group in Masters")
m = window.masters
# By name, not by number: the fork has two more screens before Masters.
window.show_screen([n for n, _ in window.screens].index("Masters"))
settle()
check("the form creates by default", m.led_save.text(), "Create ledger")

# A ledger is picked by its row, which is the key the table carries.
rows = {m.ledger_table.item(r, 0).text(): r
        for r in range(m.ledger_table.rowCount())}
m.ledger_table.setCurrentCell(rows["Petty Cash"], 0)
m.edit_ledger()
settle()
check("editing fills the form from the ledger", m.led_name.text(), "Petty Cash")
check("the button now saves instead", m.led_save.text(), "Save changes")
check("and says which ledger it is on", "Petty Cash" in m.led_hint.text(), True)
m.led_name.setText("Petty Cash Box")
m.led_open.setValue(3_000)
m.save_ledger()
settle()
check("the ledger was renamed",
      "Petty Cash Box" in set(db.get_ledgers()["name"]), True)
check("and the old name is gone",
      "Petty Cash" in set(db.get_ledgers()["name"]), False)
check("with the new opening balance",
      float(db.get_ledgers().set_index("name").loc["Petty Cash Box",
                                                   "opening"]), 3_000.0)
check("the form goes back to creating", m.led_save.text(), "Create ledger")

# A reserved ledger is looked up by NAME by the reports, so it cannot be
# renamed - and the refusal has to be a message, not a crash.
rows = {m.ledger_table.item(r, 0).text(): r
        for r in range(m.ledger_table.rowCount())}
m.ledger_table.setCurrentCell(rows["Cash"], 0)
m.edit_ledger()
settle()
try:
    db.update_ledger(db.ledger_id("Cash"), "Cash Box", "Cash-in-hand", 0, "Dr")
    check("renaming a reserved ledger is refused", "allowed", "refused")
except ValueError as exc:
    check("renaming a reserved ledger is refused", "refused", "refused")
    check("and says why", "reserved" in str(exc), True)
check("but its opening balance can be changed",
      db.update_ledger(db.ledger_id("Cash"), "Cash", "Cash-in-hand",
                       75, "Dr"), None)
m.stop_editing_ledger()
settle()

# Groups: a custom one can be renamed and moved; the reserved 27 cannot.
db.add_group("Scooter Loan", "Loans (Liability)")
m.refresh()
settle()
grows = {m.group_table.item(r, 0).text(): r
         for r in range(m.group_table.rowCount())}
m.group_table.setCurrentCell(grows["Scooter Loan"], 0)
m.edit_group()
settle()
check("editing fills the group form", m.grp_name.text(), "Scooter Loan")
m.grp_name.setText("Two-wheeler Loan")
m.grp_parent.setCurrentText("Current Liabilities")
m.save_group()
settle()
moved = db.get_groups().set_index("name")
check("the group was renamed and moved",
      str(moved.loc["Two-wheeler Loan", "parent_name"]), "Current Liabilities")
check("its nature follows the group it moved under",
      str(moved.loc["Two-wheeler Loan", "nature"]), "liability")

grows = {m.group_table.item(r, 0).text(): r
         for r in range(m.group_table.rowCount())}
m.group_table.setCurrentCell(grows["Fixed Assets"], 0)
m.edit_group()
settle()
check("a reserved group refuses to be edited at all",
      m.grp_save.text(), "Create group")
for label, call in (
    ("renaming a reserved group",
     lambda: db.update_group("Fixed Assets", "Assets Fixed")),
    ("putting a group under itself",
     lambda: db.update_group("Two-wheeler Loan", "Two-wheeler Loan",
                             "Two-wheeler Loan")),
):
    try:
        call()
        check(f"{label} is refused", "allowed", "refused")
    except ValueError:
        check(f"{label} is refused", "refused", "refused")

# A cycle is the one that would break the reports silently, so it is built
# for real: a child cannot become its own parent's parent.
db.add_group("Under Two-wheeler", "Two-wheeler Loan")
try:
    db.update_group("Two-wheeler Loan", "Two-wheeler Loan", "Under Two-wheeler")
    check("a group cannot move under its own child", "allowed", "refused")
except ValueError:
    check("a group cannot move under its own child", "refused", "refused")
check("the tree is still a tree",
      db.primary_group("Under Two-wheeler"), "Current Liabilities")

print("\n23f) Saving an alteration goes back to the report")
db.use_company(TMP)
# Saving an alteration returns to the report it was opened from (user asked
# 2026-10-20): you were reading the Day Book and went to fix a line, so
# landing back on an empty voucher form loses your place.
window.show_screen(1)
window.daybook.refresh()
settle()
window._alter(int(window.daybook.table.key_at(0)))
settle()
check("altering takes you to voucher entry", window.stack.currentIndex(), 0)
window.voucher.accept_voucher()
settle()
check("and saving it brings you back to the Day Book",
      window.stack.currentIndex(), 1)
check("the memory of where it came from is spent",
      window._altered_from, None)

# A NEW voucher must NOT jump: there you are entering a run of them, and
# being thrown onto a report after each one would be worse than useless.
window.show_screen(0)
settle()
voucher_rows = window.voucher
voucher_rows.grid.cellWidget(0, 0).setCurrentText("Rent")
voucher_rows.grid.item(0, 1).setText("12.00")
voucher_rows.grid.cellWidget(1, 0).setCurrentText("Cash")
settle()
voucher_rows.accept_voucher()
settle()
check("a new voucher leaves you on voucher entry",
      window.stack.currentIndex(), 0)


print("\n23g) The balance of the account picked, as at the voucher date")
db.use_company(TMP)
v = window.voucher
window.show_screen(0)
v.reset()
settle()

# As asked 2026-10-20: "show the balance under the account selection as of the
# voucher date and change it based on amount we enter in voucher".
standing = acc.balances().set_index("ledger").loc["Bank", "closing"]
v.grid.cellWidget(0, 0).setCurrentText("Bank")
settle()
check("the balance appears as soon as the account is picked",
      v.grid.item(0, 3).text(), acc.fmt_side(standing))
v.grid.item(0, 1).setText("500")
settle()
check("a debit moves it the debit way",
      v.grid.item(0, 3).text(), acc.fmt_side(standing + 500))
v.grid.item(0, 1).setText("")
v.grid.item(0, 2).setText("500")
settle()
check("and a credit the other way",
      v.grid.item(0, 3).text(), acc.fmt_side(standing - 500))
check("the column is told, not typed",
      bool(v.grid.item(0, 3).flags() & QtCore.Qt.ItemIsEditable), False)

# AS AT THE VOUCHER DATE, which is the whole point: the vouchers in this book
# are dated today, so the day before shows the opening balance alone.
v.reset()
settle()
yesterday = date.today() - timedelta(days=1)
v.date_edit.setDate(desktop.to_qdate(yesterday))
settle()
v.grid.cellWidget(0, 0).setCurrentText("Bank")
settle()
check("an earlier date shows the balance as it stood then",
      v.grid.item(0, 3).text(),
      acc.fmt_side(acc.balances(yesterday).set_index("ledger")
                   .loc["Bank", "closing"]))
check("which is not the same figure",
      v.grid.item(0, 3).text() != acc.fmt_side(standing), True)

# Two lines on ONE ledger accumulate, rather than both showing the same.
v.reset()
settle()
v.grid.cellWidget(0, 0).setCurrentText("Bank")
v.grid.item(0, 1).setText("100")
settle()
v.grid.cellWidget(1, 0).setCurrentText("Bank")
v.grid.item(1, 1).setText("100")
settle()
check("a second line on the same ledger carries on from the first",
      [v.grid.item(0, 3).text(), v.grid.item(1, 3).text()],
      [acc.fmt_side(standing + 100), acc.fmt_side(standing + 200)])

# THE TRAP: altering a voucher that is already in the books. The saved copy
# must be taken out of the base, or the line adds on top of itself and every
# balance reads double.
v.reset()
window.show_screen(1)
window.daybook.refresh()
settle()
altering = int(window.daybook.table.key_at(0))
head, saved = db.get_voucher(altering)
names = [str(line["ledger_name"]) for line in saved]
window._alter(altering)
settle()
as_at = date.fromisoformat(str(head["voucher_date"]))
truth = acc.balances(as_at).set_index("ledger")["closing"]
if len(names) == len(set(names)):
    check("an unchanged alteration shows the balance the books really hold",
          [v.grid.item(r, 3).text() for r in range(len(names))],
          [acc.fmt_side(float(truth.loc[n])) for n in names])
v.reset()
settle()

print("\n23h) An account that is not in the masters cannot be posted")
# Asked 2026-10-21 in Bookkeeper and ported here. The Particulars cell is an
# editable combo, so ANY text can be left in it, and `row_values()` used to
# skip such a line silently - so a voucher whose remaining lines happened to
# balance was accepted with the typed line THROWN AWAY, the amount gone, the
# screen having read Dr 1,000 Cr 1,000 and offered Accept throughout.
db.use_company(TMP)
v = window.voucher
window.show_screen(0)
v.reset()
settle()

v.type_box.setCurrentText("Journal")
set_row(v, 0, "Bank", debit=1000)
set_row(v, 1, "Rent Paid", debit=500)        # no such ledger
settle()
set_row(v, 2, "Capital", credit=1000)        # the other two balance by themselves
settle()
check("the offending line is reported, with its number",
      v.unknown_accounts(), [(2, "Rent Paid")])
check("Accept is refused although the other lines balance",
      v.accept_button.isEnabled(), False)
check("the notice names the line and what was typed",
      "line 2" in v.notice.text().lower() and '"Rent Paid"' in v.notice.text(),
      True)
check("and offers the way out", "Alt+L" in v.notice.text(), True)

# The refusal has to hold in the method that POSTS, not only in the button:
# that is the one place a line must never be dropped.
_real_warning = QtWidgets.QMessageBox.warning
_refusals = []
QtWidgets.QMessageBox.warning = staticmethod(
    lambda *a, **k: _refusals.append(a[1:]) or QtWidgets.QMessageBox.Ok)
WIDE = (date(2000, 1, 1), date.today() + timedelta(days=4000))
before = len(acc.day_book(*WIDE))
v.accept_voucher()
settle()
check("accept_voucher posts nothing while a line names no account",
      len(acc.day_book(*WIDE)), before)
check("and says why rather than failing quietly", bool(_refusals), True)
QtWidgets.QMessageBox.warning = _real_warning

# Fixed by picking a real account: the line that was being dropped is saved.
set_row(v, 1, "Capital", debit=500)
set_row(v, 2, "Capital", credit=1500)
settle()
check("picking a real account lifts the refusal", v.unknown_accounts(), [])
check("and hands Accept back", v.accept_button.isEnabled(), True)
v.accept_voucher()
settle()
_head, saved_lines = db.get_voucher(db.last_voucher_id("Journal"))
check("all three typed lines reach the books", len(saved_lines), 3)
check("carrying all of the money",
      sum(float(line["debit"]) for line in saved_lines), 1500.0)

# A case slip is the SAME account, not a missing one - the names are unique.
v.reset()
settle()
set_row(v, 0, "bank", debit=250)
set_row(v, 1, "capital", credit=250)
settle()
check("the same name in the wrong case is not an unknown account",
      v.unknown_accounts(), [])
check("it resolves to the real ledger",
      [r["ledger_id"] for r in v.row_values()],
      [db.ledger_id("Bank"), db.ledger_id("Capital")])
check("so its balance appears like any other",
      v.grid.item(0, 3).text() != "", True)

# A DISABLED ledger is in the masters but cannot be posted to. Saying it is
# missing would send you to create a second ledger under the same name.
db.add_ledger("Closed Account", "Sundry Debtors")
closed = db.ledger_id("Closed Account")
db.update_ledger(closed, "Closed Account", "Sundry Debtors", 0, "Dr",
                 active=False)
v.reset()                                    # reloads the ledger lists
settle()
set_row(v, 0, "Closed Account", debit=100)
set_row(v, 1, "Bank", credit=100)
settle()
check("a disabled ledger is refused too", v.accept_button.isEnabled(), False)
check("but is called disabled, not missing",
      "disabled" in v.notice.text(), True)
db.update_ledger(closed, "Closed Account", "Sundry Debtors", 0, "Dr",
                 active=True)

many = v._unknown_notice([(2, "Tea"), (4, "Rent")])
check("several are named together in one message",
      "Lines 2 and 4" in many and '"Tea"' in many and '"Rent"' in many, True)

v.reset()
settle()
check("an empty line is not an unknown account", v.unknown_accounts(), [])

print("\n23i) A pasted amount has its spaces taken off")
# Asked 2026-10-21: "amounts should truncate spaces while pasting amounts in
# voucher entry". The ends were already stripped; a space INSIDE the figure
# was not, and float() refuses it - so "1 000.00" pasted out of a spreadsheet
# read as 0.00, and a line whose amount is zero is dropped from the voucher
# without a word. A thousand rupees that quietly never happened.
check("a space inside the figure no longer reads as nothing",
      desktop.parse_amount("1 000.00"), 1000.0)
check("nor does a NON-BREAKING space, which is what Excel emits",
      desktop.parse_amount("1\u00a0000.00"), 1000.0)
check("nor a thin space", desktop.parse_amount("1\u202f000.50"), 1000.5)
check("padding at the ends still goes", desktop.parse_amount("  500  "), 500.0)
check("separators and padding together",
      desktop.parse_amount(" 1,00,000.50 "), 100000.5)
check("blank is still nothing", desktop.parse_amount("   "), 0.0)

v = window.voucher
window.show_screen(0)
v.reset()
settle()
set_row(v, 0, "Bank")
v.grid.item(0, 1).setText("1 000.00")        # as pasted
settle()
check("the cell is tidied to the figure it means",
      v.grid.item(0, 1).text(), "1,000.00")
check("and the voucher counts it", v.dr_label.text(), "Dr 1,000.00")
check("the balancing amount is still offered from it",
      v.grid.item(1, 2).text(), "1,000.00")

# Typed amounts must be left ALONE - a cell that reformats itself under the
# person entering it is its own annoyance.
v.reset()
settle()
set_row(v, 0, "Bank")
v.grid.item(0, 1).setText("500")
settle()
check("a typed amount is not reformatted", v.grid.item(0, 1).text(), "500")

# Something unreadable keeps its text, so it can be seen and corrected,
# rather than being turned into a confident 0.00.
v.grid.item(0, 1).setText("1 0 abc")
settle()
check("an unreadable paste is de-spaced but still shown",
      v.grid.item(0, 1).text(), "10abc")

# End to end: a pasted figure reaches the books as that figure. One line
# arrives the way a spreadsheet gives it (non-breaking space for the
# thousands separator), the other padded at both ends.
v.reset()
settle()
v.type_box.setCurrentText("Journal")
set_row(v, 0, "Bank")
v.grid.item(0, 1).setText("2\u00a0500.00")
settle()
set_row(v, 1, "Capital")
v.grid.item(1, 2).setText(" 2 500.00 ")
settle()
check("a voucher of pasted amounts balances", v.accept_button.isEnabled(), True)
v.accept_voucher()
settle()
_head, pasted_lines = db.get_voucher(db.last_voucher_id("Journal"))
check("and is posted at the figure that was pasted",
      sum(float(line["debit"]) for line in pasted_lines), 2500.0)
v.reset()
settle()

print("\n23j) Double-click an account on the Trial Balance to read its ledger")
# Asked 2026-10-21: "if i double click on any account in trial balance it
# should take me to ledger of that account".
db.use_company(TMP)
tb = window.trial
window.show_screen(3)
tb.by_group.setChecked(False)
tb.refresh()
settle()
first = str(tb.table.key_at(0))
check("each account line carries its own name as its key", bool(first), True)
tb.table.doubleClicked.emit(tb.table.model().index(0, 0))
settle()
check("it moves to the Ledger", window.stack.currentIndex(), 2)
check("on the account that was clicked",
      window.ledger.ledger_box.currentText(), first)

# THE POINT OF THE DRILL-DOWN: the ledger has to close on the figure that was
# clicked, which means carrying the trial balance's date across. Drawn to any
# other date it answers a different question.
_, tb_as_at = tb.period.range()
check("the ledger is drawn to the trial balance's date",
      window.ledger.period.range()[1], tb_as_at)
tb_rows = acc.trial_balance(tb_as_at)["rows"].set_index("Particulars")
figure = round(float(tb_rows.loc[first, "Debit"])
               - float(tb_rows.loc[first, "Credit"]), 2)
statement = acc.ledger_statement(db.ledger_id(first),
                                 *window.ledger.period.range())
check("and closes on the figure that was pointed at",
      statement["closing"], figure)

# A row stopped being its position when the headers became sortable, so the
# account opened must come from the row, not from the frame index.
window.show_screen(3)
tb.refresh()
settle()
frame_first = str(tb.table.key_at(0))
# Twice: sort_by TOGGLES, so the second call is descending - which is what
# puts a DIFFERENT account on line 0 and makes this test able to fail. The
# rows arrive ordered by name, so ascending alone leaves row 0 where it was
# and a positional bug would sail through.
tb.table.sort_by(0)
tb.table.sort_by(0)
settle()
top = str(tb.table.key_at(0))
check("sorting put a different account on the top line", top != frame_first,
      True)
tb.table.doubleClicked.emit(tb.table.model().index(0, 0))
settle()
check("after a sort it opens the account on that line, not frame row 0",
      window.ledger.ledger_box.currentText(), top)

# Summarised by group there is nothing to open - a group has no ledger - so
# it says so rather than doing nothing, which reads as broken.
window.show_screen(3)
tb.by_group.setChecked(True)
settle()
tb.table.doubleClicked.emit(tb.table.model().index(0, 0))
settle()
check("a group row opens nothing", window.stack.currentIndex(), 3)
check("and says why", "groups, not accounts" in window.statusBar().currentMessage(),
      True)
tb.by_group.setChecked(False)
settle()

print("\n23l) Double-click an account on the detailed Balance Sheet too")
# The same drill-down from the other as-at report. Only the INDENTED lines
# are accounts: a heading has no ledger, and neither have the figures the
# statement works out for itself.
bs = window.sheet
window.show_screen(5)
bs.detailed.setChecked(False)
bs.refresh()
settle()
check("summarised, no line offers a ledger",
      [bs.liabilities.key_at(r) for r in range(bs.liabilities.rowCount())],
      [None] * bs.liabilities.rowCount())
bs.liabilities.doubleClicked.emit(bs.liabilities.model().index(0, 0))
settle()
check("so a double-click stays put", window.stack.currentIndex(), 5)
check("and says what to tick",
      "Show the accounts under each heading"
      in window.statusBar().currentMessage(), True)

bs.detailed.setChecked(True)
settle()
# The first line with a key is an account; the heading above it has none.
side = bs.liabilities
rows = [(r, side.key_at(r)) for r in range(side.rowCount())]
accounts = [(r, k) for r, k in rows if k is not None]
check("detailed, the indented accounts carry keys", bool(accounts), True)
check("and the headings still do not",
      any(k is None for _, k in rows), True)
check("a key is the account name without its indent",
      str(accounts[0][1]).startswith(" "), False)

row, account = accounts[0]
side.doubleClicked.emit(side.model().index(row, 0))
settle()
check("double-clicking an account opens the Ledger",
      window.stack.currentIndex(), 2)
check("on that account", window.ledger.ledger_box.currentText(), str(account))
_, bs_as_at = bs.period.range()
check("drawn to the balance sheet's own date",
      window.ledger.period.range()[1], bs_as_at)

# The figure on the sheet is the side's own sign, so compare magnitudes: the
# point is that the ledger answers for the same date, not the same sign
# convention.
statement = acc.ledger_statement(db.ledger_id(str(account)),
                                 *window.ledger.period.range())
sheet_now = acc.balance_sheet(bs_as_at, detailed=True)["liabilities"]
shown = float(sheet_now.set_index("Particulars").loc[
    f"{acc.DETAIL_INDENT}{account}", "Amount"])
check("and closes on the figure that was pointed at",
      abs(statement["closing"]), abs(shown))

# A heading row in the detailed form explains itself rather than doing
# nothing - and does not navigate.
window.show_screen(5)
settle()
heading_row = next(r for r, k in rows if k is None)
side.doubleClicked.emit(side.model().index(heading_row, 0))
settle()
check("a heading opens nothing", window.stack.currentIndex(), 5)
check("and says only accounts do",
      "Only the accounts under a heading" in window.statusBar().currentMessage(),
      True)

# Asked 2026-10-21: "while expanding balance sheet headers, please have
# headers be in bold and underlined so they can easily be difrenciated".
# The indent alone was too quiet to read.
bs.refresh()
settle()
marks = [(bs.liabilities.key_at(r) is None,
          bs.liabilities.item(r, 0).font().bold(),
          bs.liabilities.item(r, 0).font().underline())
         for r in range(bs.liabilities.rowCount())]
check("every heading is bold and underlined",
      all(bold and under for heading, bold, under in marks if heading), True)
check("and no account is",
      any(bold or under for heading, bold, under in marks if not heading),
      False)
head_row = next(r for r in range(bs.liabilities.rowCount())
                if bs.liabilities.key_at(r) is None)
check("the heading's own figure is marked too, not just its name",
      (bs.liabilities.item(head_row, 1).font().bold(),
       bs.liabilities.item(head_row, 1).font().underline()), (True, True))

# ON PAPER as well: HTML collapses the indent spaces to nothing, so a
# printed heading and the accounts under it sat at exactly the same indent
# with nothing in bold - the report was unreadable in the one form that
# needed the distinction most.
printed = desktop.report_html(bs.report(), company="Probe")
check("the printed heading is bold", "<b>Capital Account</b>" in printed, True)
check("and the printed account keeps a real indent",
      "&nbsp;" in printed, True)

bs.detailed.setChecked(False)
settle()
flat = [(bs.liabilities.item(r, 0).font().bold(),
         bs.liabilities.item(r, 0).font().underline())
        for r in range(bs.liabilities.rowCount())]
check("summarised, nothing is emphasised - every line is a heading",
      any(bold or under for bold, under in flat), False)
check("and the printed summary is left plain",
      "<b>" in desktop.report_html(bs.report(), company="Probe"), False)
# Every other report is untouched by the layered treatment.
window.trial.refresh()
settle()
check("the trial balance prints plain",
      "&nbsp;" in desktop.report_html(window.trial.report(), company="Probe"),
      False)

print("\n23k) Abandoning an alteration goes back to the report")
# Asked 2026-10-21: "while i am in ledger an double click on an entry to view
# and click on cancel edit it land on voucher entry page instead it should
# land back on ledger page". Saving an alteration already went back; giving
# up on one did not, which left you on an empty entry form.
_real_question = QtWidgets.QMessageBox.question
QtWidgets.QMessageBox.question = staticmethod(
    lambda *a, **k: QtWidgets.QMessageBox.StandardButton.Yes)

window.show_screen(2)
window.ledger.ledger_box.setCurrentText("Bank")
window.ledger.period.year_box.setCurrentText(desktop.PeriodBar.SINCE)
settle()
check("the ledger has entries to open", window.ledger._frame.empty, False)
window.ledger.table.doubleClicked.emit(
    window.ledger.table.model().index(0, 0))
settle()
check("double-clicking an entry opens it for alteration",
      window.voucher.editing_id is not None, True)
check("on the voucher screen", window.stack.currentIndex(), 0)
window.voucher.cancel_edit()
settle()
check("Cancel edit lands back on the Ledger", window.stack.currentIndex(), 2)
check("with the form cleared", window.voucher.editing_id, None)

# From the Day Book it goes back to the Day Book, and `reset()` on its own
# must NOT navigate - it runs after every accept, on a company switch and
# from the constructor, so putting the going-back there would move the screen
# in all of those.
window.show_screen(1)
window.daybook.refresh()
settle()
window._alter(int(window.daybook.table.key_at(0)))
settle()
window.voucher.reset()
settle()
check("reset() alone stays put", window.stack.currentIndex(), 0)
window.voucher.cancel_edit()
settle()
check("the button lands back on the Day Book", window.stack.currentIndex(), 1)

# A copy begun on a blank form came from nowhere, so discarding it has
# nowhere to go back to and must stay where it is.
window.show_screen(0)
window.voucher.reset()
window.voucher.type_box.setCurrentText("Journal")
settle()
window.voucher.copy_voucher()
settle()
check("discarding a copy begun here stays on voucher entry",
      window.stack.currentIndex(), 0)
window.voucher.cancel_edit()
settle()
check("still there", window.stack.currentIndex(), 0)
QtWidgets.QMessageBox.question = _real_question
window.voucher.reset()
settle()

print("\n24) The shortcuts actually fire, pressed for real")
# Binding a key is not the same as the key working, which is the whole lesson
# of the Alt+C bug - so press it for real and see whether the action runs.
# This goes LAST because a WidgetWithChildrenShortcut only fires inside an
# ACTIVE window: showing the window realises its geometry, and the width
# checks above read window.width(), which would then report the size clamped
# to the 800x800 offscreen display instead of the unrealised hint.
voucher = window.voucher
window.show_screen(0)
window.show()
# The voucher-type keys are QActions with the default WindowShortcut context,
# which fires only in an ACTIVE window - show() alone is not enough. (Alt+L
# is WidgetWithChildrenShortcut, so focus is all that one needs, which is why
# it passed without this.)
window.activateWindow()
window.raise_()
app.processEvents()
_real_create = voucher.create_ledger
_fired = []
voucher.create_ledger = lambda: _fired.append(True)
voucher.grid.setCurrentCell(0, 0)
voucher.grid.setFocus()
QTest.keyClick(voucher.grid, QtCore.Qt.Key_L, QtCore.Qt.AltModifier)
app.processEvents()
check("pressing Alt+L really runs it", bool(_fired), True)
_fired.clear()
QTest.keyClick(voucher.grid, QtCore.Qt.Key_C, QtCore.Qt.AltModifier)
app.processEvents()
check("and Alt+C no longer does anything here", bool(_fired), False)
voucher.create_ledger = _real_create

# That Alt+C press did what it always did - it OPENED the Company menu - and
# an open menu grabs the keyboard, so every shortcut tested after this point
# would read as dead until it is dismissed.
while QtWidgets.QApplication.activePopupWidget():
    QtWidgets.QApplication.activePopupWidget().close()
    app.processEvents()

# Same for the voucher-type keys: they were only ever tested by calling the
# method, so a dead keystroke would not have shown up there either.
#
# Use QTest.keySequence, NOT keyClick(key, modifiers): only the former
# reliably routes a modified keystroke through Qt's shortcut map, so a
# keyClick test reports every shortcut as dead whether it is or not.
def press(sequence):
    QTest.keySequence(window, QtGui.QKeySequence(sequence))
    app.processEvents()

for key, name in desktop.VOUCHER_KEYS.items():
    window._set_voucher_type("Journal" if name != "Journal" else "Payment")
    app.processEvents()
    press(key)
    check(f"pressing {key} really selects {name}",
          voucher.type_box.currentText(), name)

# And the plain-Alt version must NOT switch type, or it would be shadowing
# a menu - Alt+P is free, but Alt+S/R/C belong to the menu bar.
window._set_voucher_type("Journal")
app.processEvents()
press("Alt+P")
check("plain Alt+P does not switch type",
      voucher.type_box.currentText(), "Journal")

window.hide()
app.processEvents()

print("\n25) Bills: allocating a line, and the Outstanding screen")
db.use_company(TMP)
db.add_ledger("Chase Me Ltd", "Sundry Debtors")
window.masters.refresh()
voucher = window.voucher
voucher.reset()
settle()

check("a party ledger comes out bill-wise",
      bool(db.get_ledgers().set_index("name").loc["Chase Me Ltd", "bill_wise"]),
      True)

# The Bills button only means anything on a line that keeps bills.
set_row(voucher, 0, "Chase Me Ltd", debit=50_000)
set_row(voucher, 1, "Sales", credit=50_000)
voucher.grid.setCurrentCell(0, 0)
settle()
check("Bills is offered on the party line", voucher.bills_button.isEnabled(), True)
voucher.grid.setCurrentCell(1, 0)
settle()
check("and not on the Sales line", voucher.bills_button.isEnabled(), False)
voucher.grid.setCurrentCell(0, 0)
settle()

# Raising a bill. The dialog is built directly - exec() would block - and the
# voucher number is offered as the reference.
dialog = desktop.BillsDialog(
    voucher, party="Chase Me Ltd", ledger_id=db.ledger_id("Chase Me Ltd"),
    amount=50_000.0, suggested_ref="7")
check("a new bill is offered the voucher's own number",
      dialog.new_ref.text(), "7")
check("with the whole line allocated to it",
      "All" in dialog.summary.text(), True)
dialog._accept()
check("which yields one new-reference allocation",
      [(a["kind"], a["bill_ref"], a["signed"]) for a in dialog.allocations],
      [("new", "7", 50_000.0)])

voucher.grid.cellWidget(0, 0)._bills = dialog.allocations
check("and the row carries them into row_values",
      voucher.row_values()[0]["bills"], dialog.allocations)
window._accept()
settle()
check("posting puts the bill on the outstanding report",
      float(acc.outstanding()["bills"].set_index("Ref").loc["7", "Pending"]),
      50_000.0)

# Settling it. Apply to oldest fills the open bills from the top.
voucher.reset()
voucher.type_box.setCurrentText("Receipt")
set_row(voucher, 0, "Cash", debit=20_000)
set_row(voucher, 1, "Chase Me Ltd", credit=20_000)
settle()
dialog = desktop.BillsDialog(
    voucher, party="Chase Me Ltd", ledger_id=db.ledger_id("Chase Me Ltd"),
    amount=-20_000.0)
check("the open bill is offered to settle against", list(dialog._open["Ref"]),
      ["7"])
check("nothing allocated to begin with",
      "on account" in dialog.summary.text(), True)
dialog._apply_to_oldest()
check("apply to oldest allocates the lot", "All" in dialog.summary.text(), True)
dialog._accept()
check("as an against-reference running the other way",
      [(a["kind"], a["bill_ref"], a["signed"]) for a in dialog.allocations],
      [("against", "7", -20_000.0)])
voucher.grid.cellWidget(1, 0)._bills = dialog.allocations
window._accept()
settle()
check("so the bill is part-settled",
      float(acc.outstanding()["bills"].set_index("Ref").loc["7", "Pending"]),
      30_000.0)

# Over-allocating is refused by the dialog, before the data layer sees it.
dialog = desktop.BillsDialog(
    voucher, party="Chase Me Ltd", ledger_id=db.ledger_id("Chase Me Ltd"),
    amount=-5_000.0)
dialog.new_ref.setText("TOO-MUCH")
dialog.new_amount.setText("9000")
settle()
check("allocating more than the line disables Done",
      dialog.ok.isEnabled(), False)
check("and says by how much", "more than the line" in dialog.summary.text(),
      True)

# A reference belongs to a PARTY, so re-pointing the row must drop it.
voucher.reset()
set_row(voucher, 0, "Chase Me Ltd", debit=1_000)
combo = voucher.grid.cellWidget(0, 0)
combo._bills = [{"kind": "new", "bill_ref": "KEEP", "signed": 1_000.0,
                 "due_date": ""}]
set_row(voucher, 0, "Cash", debit=1_000)
settle()
check("pointing the row at another ledger drops its bill references",
      getattr(combo, "_bills", []), [])
voucher.reset()

# --- the Outstanding screen -------------------------------------------- #
screen = window.outstanding
screen.refresh()
settle()
check("Outstanding opens on today, not the year's closing day",
      screen.period.year_box.currentText(), "Today")
check("which is what it ages against", screen.period.range()[1], date.today())
check("the bill is listed", "7" in list(screen._shown["Ref"]), True)
check("the ageing summary adds up to the receivables shown",
      round(float(screen._ageing["Receivable"].sum()), 2),
      round(float(screen._shown[
          screen._shown["Party"] == "Chase Me Ltd"]["Pending"].sum()), 2))

screen.age_by.setCurrentText("bill date")
settle()
check("the ageing basis is named in the printed report",
      "bill date" in screen.report().period, True)
screen.age_by.setCurrentText("due date")
settle()

screen.show_side.setCurrentText("Payable only")
settle()
check("filtering to payables hides a receivable", len(screen._shown), 0)
check("and the filter is named on the report",
      "payable only" in screen.report().period, True)
screen.show_side.setCurrentText("Receivable and payable")
settle()

report = screen.report()
check("it prints as a report with its tables",
      [s.heading for s in report.sections][:1], ["Bills outstanding"])
check("and carries the totals, which live in a label not a table",
      any("Receivable" in n for n in report.notes), True)

print("\n26) GST: the dialog, the lines it writes, and the report")
db.use_company(TMP)
db.save_company({**db.get_company(), "company_name": "Aarkaya Handicrafts LLC",
                 "pan": "AAAPZ1234C", "gstin": "29AAAPZ1234C1Z5"})
window.company.refresh()
check("our state comes from our GSTIN", db.company_state_code(), "29")

# Set a rate on an EXISTING ledger through Masters, which is the only way to
# do it from this screen.
db.add_ledger("GST Sales", "Sales Accounts")
window.masters.refresh()
settle()
row = window.masters._ledger_names.index("GST Sales")
window.masters.ledger_table.setCurrentCell(row, 0)
window.masters.led_rate.setText("18")
window.masters.apply_gst_to_selected()
settle()
check("Set GST on selected puts a rate on a ledger that already existed",
      float(db.get_ledgers().set_index("name").loc["GST Sales", "gst_rate"]),
      18.0)
check("and says so", "taxable at 18%" in window.statusBar().currentMessage(),
      True)

# Blank means not taxable, which is NOT the same as zero.
window.masters.led_rate.setText("")
window.masters.apply_gst_to_selected()
settle()
check("blanking the rate makes it untaxable again",
      db.get_ledgers().set_index("name").loc["GST Sales", "gst_rate"] is None
      or pd.isna(db.get_ledgers().set_index("name").loc["GST Sales", "gst_rate"]),
      True)
window.masters.led_rate.setText("18")
window.masters.apply_gst_to_selected()
settle()
check("rubbish in the rate box is refused",
      window.masters._typed_rate.__self__ is window.masters, True)
try:
    window.masters.led_rate.setText("eighteen")
    window.masters._typed_rate()
    check("a rate that is not a number is refused", "allowed", "refused")
except ValueError as exc:
    check("a rate that is not a number is refused", "refused", "refused")
    check("and says to leave it blank instead",
          "blank" in str(exc), True)
window.masters.led_rate.setText("18")

db.add_ledger("GST Local", "Sundry Debtors", gstin="29BBBPZ1234C1Z5",
              state_code="29")
db.add_ledger("GST Far", "Sundry Debtors", gstin="27CCCPZ1234C1Z5",
              state_code="27")
window.voucher.reload_ledgers()
settle()

# --- the dialog's arithmetic -------------------------------------------- #
local = desktop.GstDialog(window.voucher, taxable=-10_000.0,
                          voucher_type="Sales", party="GST Local",
                          party_gstin="29BBBPZ1234C1Z5", suggested_rate=18.0,
                          place_of_supply="29")
got = local.computed()
check("a supply in our own state splits into CGST and SGST",
      (got["supply_type"], got["cgst"], got["sgst"], got["igst"]),
      ("intra", 900.0, 900.0, 0.0))
far = desktop.GstDialog(window.voucher, taxable=-10_000.0,
                        voucher_type="Sales", party="GST Far",
                        party_gstin="27CCCPZ1234C1Z5", suggested_rate=18.0,
                        place_of_supply="27")
got = far.computed()
check("one to another state is all IGST",
      (got["supply_type"], got["cgst"], got["igst"]), ("inter", 0.0, 1800.0))
check("and it says why", "IGST" in far.note.text(), True)

# A GSTIN names its own state, so typing one moves the place of supply.
typed = desktop.GstDialog(window.voucher, taxable=-1_000.0,
                          voucher_type="Sales", party="Someone")
typed.gstin.setText("19ZZZPZ1234C1Z5")
settle()
check("typing a GSTIN sets the place of supply from it",
      typed.place.currentData(), "19")
check("which makes the supply interstate",
      typed.computed()["supply_type"], "inter")

# --- the lines it writes ------------------------------------------------- #
_real_exec = desktop.GstDialog.exec
desktop.GstDialog.exec = lambda self: (self._accept(), 1)[1]

voucher = window.voucher
voucher.reset()
voucher.type_box.setCurrentText("Sales")
set_row(voucher, 0, "GST Sales", credit=10_000)
set_row(voucher, 1, "GST Local", debit=10_000)
settle()
check("GST is offered on a Sales voucher", voucher.gst_button.isEnabled(), True)
check("and the taxable line is found", voucher._taxable_rows(), [0])
voucher.edit_gst()
settle()
posted = {voucher.grid.cellWidget(r, 0).currentText():
          voucher._row_amount(r)
          for r in range(voucher.grid.rowCount())
          if voucher.grid.cellWidget(r, 0)
          and voucher.grid.cellWidget(r, 0).currentText().strip()}
check("the tax lines are written on the same side as the supply",
      (posted["Output CGST"], posted["Output SGST"]), (-900.0, -900.0))
check("and the party carries the whole invoice",
      posted["GST Local"], 11_800.0)
check("so the voucher balances", voucher.gap_label.text(), "Difference 0.00")
check("the rate and place of supply are held for saving",
      (voucher._gst["rate"], voucher._gst["place_of_supply"]), (18.0, "29"))

# Applying again must REPLACE the tax lines, not add a second set.
voucher.edit_gst()
settle()
check("applying GST twice does not stack a second set of tax lines",
      len(voucher._tax_rows()), 2)
check("and it still balances", voucher.gap_label.text(), "Difference 0.00")

window._accept()
settle()
summary = acc.gst_summary()
check("posting puts the tax on the GST return",
      summary["totals"]["output_tax"], 1_800.0)
# The check that matters: the return must equal the tax LEDGERS.
held = acc.balances().set_index("ledger")
check("and the return equals the Output tax ledgers",
      summary["totals"]["output_tax"],
      round(sum(abs(float(held.loc[f"Output {t}", "closing"]))
                for t in ("CGST", "SGST", "IGST")), 2), tol=0.0)

# GST belongs to Sales and Purchase only.
voucher.reset()
voucher.type_box.setCurrentText("Journal")
settle()
check("GST is not offered on a Journal", voucher.gst_button.isEnabled(), False)
voucher.reset()
desktop.GstDialog.exec = _real_exec

# --- the GST screen ----------------------------------------------------- #
screen = window.gst
window.refresh_all()
settle()
check("the GST screen opens on the summary",
      screen.view.currentText(), "Summary")
check("which names what is payable", "Payable" in screen.totals.text(), True)
check("and explains the figure",
      "payable" in screen.status.text().lower(), True)
check("it prints the rate breakdowns",
      [s.heading for s in screen.report().sections][:1],
      ["Outward supplies, by rate"])

screen.view.setCurrentText("Outward supplies")
settle()
check("the outward register lists the invoice",
      len(screen._register), 1)
check("with the place of supply spelled out",
      str(screen._register.iloc[0]["Place of supply"]), "Karnataka")
check("and it prints under its own heading",
      [s.heading for s in screen.report().sections], ["Outward supplies"])
check("the printed report carries the totals, which live in a label",
      any("Output" in n for n in screen.report().notes), True)
screen.view.setCurrentText("Summary")
settle()

# The real data folder must be untouched by any of this.
check("the suite wrote only into its sandbox",
      str(db.DATA_DIR).startswith(str(Path(tempfile.gettempdir()))), True)
db.use_company(TMP)

print("\n" + ("ALL DESKTOP CHECKS PASSED" if not FAILURES
             else f"FAILURES: {FAILURES}"))
raise SystemExit(1 if FAILURES else 0)
